"""Exporting accounts and contacts.

Row scope narrows a plain CSM to their own portfolio, and the workbook that
comes back actually contains what it claims to. An ACM may reach this page
like everyone else - export is read-only, and the accounts export already
drops the health columns for anyone without ``health.view``.
"""

from __future__ import annotations

import io

import pytest
from django.urls import reverse
from openpyxl import load_workbook

from csm.contacts.models import Contact
from tests.factories import AccountFactory, UserFactory
from tests.test_views import PASSWORD, login


@pytest.fixture
def staff(db):
    users = {
        "csm": UserFactory(csm=True, email="exp-csm@dphish.com"),
        "senior": UserFactory(senior_csm=True, email="exp-senior@dphish.com"),
        "acm": UserFactory(acm=True, email="exp-acm@dphish.com"),
        "manager": UserFactory(manager=True, email="exp-manager@dphish.com"),
    }
    for user in users.values():
        user.set_password(PASSWORD)
        user.save(update_fields=["password"])
    return users


@pytest.fixture
def portfolio(staff):
    return {
        "mine": AccountFactory(
            client_code="EXP01", client_name="Mine Co", csm=staff["csm"]
        ),
        "not_mine": AccountFactory(
            client_code="EXP02", client_name="Not Mine Co", csm=staff["senior"]
        ),
    }


def workbook_rows(content: bytes):
    sheet = load_workbook(io.BytesIO(content)).worksheets[0]
    return list(sheet.iter_rows(values_only=True))


@pytest.mark.parametrize("who", ["csm", "senior", "manager", "acm"])
def test_everyone_may_reach_the_export_page(client, staff, who):
    login(client, staff[who])
    assert client.get(reverse("importer:export")).status_code == 200


def test_an_acms_accounts_export_carries_no_health_column(client, staff, portfolio):
    login(client, staff["acm"])
    response = client.get(reverse("importer:export_download", args=["accounts"]))
    rows = workbook_rows(response.content)
    assert "Health Status" not in rows[0]
    assert "Health Score" not in rows[0]


def test_a_plain_csm_only_exports_their_own_portfolio(client, staff, portfolio):
    login(client, staff["csm"])
    response = client.get(reverse("importer:export_download", args=["accounts"]))
    rows = workbook_rows(response.content)

    codes = [row[0] for row in rows[1:]]
    assert codes == ["EXP01"]


def test_a_senior_csm_exports_the_whole_portfolio(client, staff, portfolio):
    login(client, staff["senior"])
    response = client.get(reverse("importer:export_download", args=["accounts"]))
    rows = workbook_rows(response.content)

    codes = {row[0] for row in rows[1:]}
    assert codes == {"EXP01", "EXP02"}


def test_the_accounts_export_names_the_csm_by_email(client, staff, portfolio):
    login(client, staff["senior"])
    response = client.get(reverse("importer:export_download", args=["accounts"]))
    rows = workbook_rows(response.content)

    header = list(rows[0])
    assert "CSM Email" in header
    csm_column = header.index("CSM Email")
    by_code = {row[0]: row[csm_column] for row in rows[1:]}
    assert by_code["EXP01"] == staff["csm"].email


def test_the_contacts_export_is_scoped_the_same_way(client, staff, portfolio):
    Contact.objects.create(
        account=portfolio["mine"], full_name="Dana Farouk", email="dana@acme.example"
    )
    Contact.objects.create(
        account=portfolio["not_mine"], full_name="Omar Said", email="omar@acme.example"
    )

    login(client, staff["csm"])
    response = client.get(reverse("importer:export_download", args=["contacts"]))
    rows = workbook_rows(response.content)

    names = [row[1] for row in rows[1:]]
    assert names == ["Dana Farouk"]


def test_the_contacts_export_header_matches_the_import_template(client, staff, portfolio):
    from csm.importer.contact_parsing import CONTACT_COLUMNS
    from csm.importer.template import header_label

    login(client, staff["csm"])
    response = client.get(reverse("importer:export_download", args=["contacts"]))
    header = list(workbook_rows(response.content)[0])

    expected = [header_label(field, CONTACT_COLUMNS) for field in CONTACT_COLUMNS]
    assert header == expected


def test_an_unedited_contacts_export_reimports_clean(portfolio, staff):
    """Mirrors the accounts round-trip: an unedited export re-imports with
    nothing lost and nothing flagged unknown, including the multi-select
    channels column and the focal rank."""
    from csm.contacts.models import Contact, ContactChannel
    from csm.core.enums import CommunicationChannel, Level
    from csm.importer.contact_parsing import parse_contacts
    from csm.importer.export import build_contacts_xlsx

    contact = Contact.objects.create(
        account=portfolio["mine"],
        full_name="Dana Farouk",
        title="Head of Security",
        email="dana@acme.example",
        phone="+20 100 000 0000",
        influence=Level.HIGH,
        interest=Level.MEDIUM,
        preferred_channel=CommunicationChannel.EMAIL,
        notes="Prefers async updates.",
    )
    contact.set_focal_point(True)
    contact.save(update_fields=["is_focal_point", "focal_rank"])
    ContactChannel.objects.create(contact=contact, channel=CommunicationChannel.EMAIL)
    ContactChannel.objects.create(contact=contact, channel=CommunicationChannel.WHATSAPP)

    content = build_contacts_xlsx(Contact.objects.filter(pk=contact.pk))
    parsed, unknown = parse_contacts(
        io.BytesIO(content), known_codes={portfolio["mine"].client_code.lower()},
        filename="export.xlsx",
    )

    assert unknown == []
    row = parsed[0]
    assert row.errors == []
    assert row.cleaned["title"] == "Head of Security"
    assert row.cleaned["email"] == "dana@acme.example"
    assert row.cleaned["phone"] == "+20 100 000 0000"
    assert row.cleaned["influence"] == Level.HIGH
    assert row.cleaned["interest"] == Level.MEDIUM
    assert set(row.cleaned["channels"]) == {
        CommunicationChannel.EMAIL,
        CommunicationChannel.WHATSAPP,
    }
    assert row.cleaned["preferred_channel"] == CommunicationChannel.EMAIL
    assert row.cleaned["focal_rank"] == 1
    assert row.cleaned["is_focal_point"] is True
    assert row.cleaned["notes"] == "Prefers async updates."


def test_the_download_is_a_readable_workbook(client, staff, portfolio):
    login(client, staff["senior"])
    response = client.get(reverse("importer:export_download", args=["accounts"]))

    assert response.status_code == 200
    assert "accounts-export.xlsx" in response["Content-Disposition"]
    assert response.content


def test_the_accounts_export_header_matches_the_import_template(client, staff, portfolio):
    """Same fields, same order as the importer's own template - not a subset
    of it and not reshuffled - so an edited export lines up column-for-column
    against what somebody filling in the template sees."""
    from csm.importer.parsing import COLUMNS
    from csm.importer.template import header_label

    login(client, staff["senior"])
    response = client.get(reverse("importer:export_download", args=["accounts"]))
    header = list(workbook_rows(response.content)[0])

    expected = [header_label(field) for field in COLUMNS]
    assert header[: len(expected)] == expected


def test_an_unedited_accounts_export_reimports_clean(portfolio, staff):
    """The point of matching the importer's own columns: exporting and
    re-importing an account, untouched, produces no errors and loses nothing
    the importer is able to read back."""
    import datetime as dt

    from csm.accounts.models import Account, Partner
    from csm.core.enums import (
        AccountClass,
        AccountSize,
        Deployment,
        LicenseType,
        MeetingFrequency,
        ReviewStatus,
    )
    from csm.importer.export import build_accounts_xlsx
    from csm.importer.parsing import parse

    account = portfolio["mine"]
    account.country = "EG"
    account.location = "Cairo"
    account.deployment = Deployment.ON_PREM
    account.is_managed_service = True
    account.license_type = LicenseType.COMPLETE
    account.boq_users_count = 120
    account.added_users_count = 15
    account.contract_start_date = dt.date(2024, 1, 1)
    account.contract_end_date = dt.date(2026, 12, 31)
    account.partner = Partner.objects.create(name="Nile Reseller")
    account.quarterly_meeting_date = dt.date(2026, 3, 15)
    account.events_count = 4
    account.digital_content_count = 2
    account.event_location = "Riyadh"
    account.g2_review_status = ReviewStatus.PUBLISHED
    account.gartner_review_status = ReviewStatus.REQUESTED
    account.account_class = AccountClass.A
    account.size = AccountSize.ENTERPRISE
    account.required_meeting_frequency = MeetingFrequency.WEEKLY
    account.save()

    content = build_accounts_xlsx(
        Account.objects.filter(pk=account.pk), include_health=False
    )
    people = {
        "by_email": {
            staff["csm"].email.lower(): staff["csm"],
            account.acm.email.lower(): account.acm,
        },
        "by_name": {},
    }
    parsed, unknown = parse(io.BytesIO(content), people=people, filename="export.xlsx")

    assert unknown == []
    row = parsed[0]
    assert row.errors == []
    assert row.cleaned["country"] == "EG"
    assert row.cleaned["deployment"] == Deployment.ON_PREM
    assert row.cleaned["license_type"] == LicenseType.COMPLETE
    assert row.cleaned["boq_users_count"] == 120
    assert row.cleaned["added_users_count"] == 15
    assert row.cleaned["contract_start_date"] == dt.date(2024, 1, 1)
    assert row.cleaned["contract_end_date"] == dt.date(2026, 12, 31)
    assert row.cleaned["partner_name"] == "Nile Reseller"
    assert row.cleaned["quarterly_meeting_date"] == dt.date(2026, 3, 15)
    assert row.cleaned["events_count"] == 4
    assert row.cleaned["digital_content_count"] == 2
    assert row.cleaned["event_location"] == "Riyadh"
    assert row.cleaned["g2_review_status"] == ReviewStatus.PUBLISHED
    assert row.cleaned["gartner_review_status"] == ReviewStatus.REQUESTED
    assert row.cleaned["account_class"] == AccountClass.A
    assert row.cleaned["size"] == AccountSize.ENTERPRISE
    assert row.cleaned["required_meeting_frequency"] == MeetingFrequency.WEEKLY
    assert row.cleaned["csm_id"] == staff["csm"].pk
    assert row.cleaned["acm_id"] == account.acm_id

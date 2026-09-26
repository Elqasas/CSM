"""The account sheet carries every field the account form does, in its order.

Import template, export and form are one arrangement: a person moving between
the edit page and a spreadsheet meets the same fields in the same sequence,
custom fields included, placed before Description exactly as the form places
them. The logo is the one exception - an image has no cell to live in.

The other load-bearing rule is that an update writes only the columns a sheet
actually has. A sheet exported before Description existed must not blank every
description on re-import.
"""

from __future__ import annotations

import datetime as dt
import io

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook

from csm.accounts.forms import CANDIDATE_FIELDS, CustomFieldDefinitionForm
from csm.accounts.models import (
    Account,
    AccountCustomFieldDefinition,
    AccountCustomFieldValue,
)
from csm.core.enums import CustomFieldType
from csm.importer import services
from csm.importer.export import build_accounts_xlsx
from csm.importer.models import ImportBatch
from csm.importer.parsing import COLUMNS, account_columns, parse
from tests.factories import AccountFactory, UserFactory


@pytest.fixture
def senior(db):
    return UserFactory(senior_csm=True, email="sheet-senior@dphish.com")


@pytest.fixture
def fields(db):
    return {
        "text": AccountCustomFieldDefinition.objects.create(name="Go-live owner"),
        "choice": AccountCustomFieldDefinition.objects.create(
            name="Segment",
            field_type=CustomFieldType.DROPDOWN,
            choices_text="Banking\nGovernment\nRetail",
        ),
        "when": AccountCustomFieldDefinition.objects.create(
            name="Kick-off", field_type=CustomFieldType.DATETIME
        ),
    }


def active(fields):
    return list(AccountCustomFieldDefinition.objects.filter(pk__in=[f.pk for f in fields.values()]))


def workbook(headers, rows) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def apply_sheet(content: bytes, actor) -> ImportBatch:
    batch = ImportBatch.objects.create(
        upload=SimpleUploadedFile("sheet.xlsx", content),
        original_name="sheet.xlsx",
        uploaded_by=actor,
        baseline_date=dt.date(2026, 9, 26),
    )
    batch = services.validate(batch)
    services.apply(batch, actor=actor)
    return batch


# ---------------------------------------------------------------------------
# Arrangement
# ---------------------------------------------------------------------------


def test_the_sheet_follows_the_account_form_field_for_field():
    assert list(COLUMNS) == [f for f in CANDIDATE_FIELDS if f != "logo"]


def test_custom_fields_sit_just_before_description(fields):
    columns = list(account_columns(active(fields)))

    custom = [f"custom_{d.pk}" for d in active(fields)]
    assert columns[-len(custom) - 1 : -1] == custom
    assert columns[-1] == "description"


def test_the_export_header_names_custom_fields_as_the_form_does(senior, fields):
    AccountFactory(client_code="SH01", csm=None)
    content = build_accounts_xlsx(
        Account.objects.all(), include_health=False, custom_fields=active(fields)
    )
    header = [c.value for c in load_workbook(io.BytesIO(content)).worksheets[0][1]]

    assert header[-4:] == ["Go-live owner", "Kick-off", "Segment", "Description"]
    assert "Tenant URL" in header and "LinkedIn" in header and "Renewal Count" in header


# ---------------------------------------------------------------------------
# Round trip
# ---------------------------------------------------------------------------


def test_an_unedited_export_restores_every_form_field(senior, fields):
    account = AccountFactory(
        client_code="SH02",
        csm=None,
        acm=None,
        renewal_count=3,
        website_url="https://acme.example",
        tenant_url="https://acme.tenant.example",
        drive_url="https://drive.example/acme",
        linkedin_url="https://www.linkedin.com/company/acme",
        g2_review_url="https://g2.example/acme",
        gartner_review_url="https://gartner.example/acme",
        description="Renewal runs through the partner.\nCISO is the sponsor.",
    )
    for key, value in (
        ("text", "Dana"),
        ("choice", "Government"),
        ("when", "2026-05-01T09:30:00"),
    ):
        AccountCustomFieldValue.objects.create(
            account=account, definition=fields[key], value=value
        )

    content = build_accounts_xlsx(
        Account.objects.filter(pk=account.pk),
        include_health=False,
        custom_fields=active(fields),
    )

    # Wipe everything the sheet carries, then let the export put it back.
    Account.objects.filter(pk=account.pk).update(
        renewal_count=0, website_url="", tenant_url="", drive_url="",
        linkedin_url="", g2_review_url="", gartner_review_url="", description="",
    )
    AccountCustomFieldValue.objects.filter(account=account).update(value="")

    batch = apply_sheet(content, senior)
    assert batch.unknown_columns == []
    assert list(batch.rows.values_list("errors", flat=True)) == [[]]

    account.refresh_from_db()
    assert account.renewal_count == 3
    assert account.website_url == "https://acme.example"
    assert account.tenant_url == "https://acme.tenant.example"
    assert account.drive_url == "https://drive.example/acme"
    assert account.linkedin_url == "https://www.linkedin.com/company/acme"
    assert account.g2_review_url == "https://g2.example/acme"
    assert account.gartner_review_url == "https://gartner.example/acme"
    assert account.description == "Renewal runs through the partner.\nCISO is the sponsor."
    stored = dict(
        account.custom_field_values.values_list("definition__name", "value")
    )
    assert stored == {
        "Go-live owner": "Dana",
        "Segment": "Government",
        "Kick-off": "2026-05-01T09:30:00",
    }


# ---------------------------------------------------------------------------
# Absent versus empty
# ---------------------------------------------------------------------------


def test_a_sheet_without_a_column_leaves_that_field_alone(senior, fields):
    account = AccountFactory(
        client_code="SH03",
        location="Cairo",
        description="Hand-written notes.",
        tenant_url="https://t.example",
        renewal_count=2,
    )
    AccountCustomFieldValue.objects.create(
        account=account, definition=fields["text"], value="Dana"
    )

    apply_sheet(
        workbook(["Client Code", "Client Name", "Location"], [["SH03", "Renamed", "Giza"]]),
        senior,
    )

    account.refresh_from_db()
    assert (account.client_name, account.location) == ("Renamed", "Giza")
    assert account.description == "Hand-written notes."
    assert account.tenant_url == "https://t.example"
    assert account.renewal_count == 2
    assert account.custom_field_values.get().value == "Dana"


def test_an_empty_cell_in_a_present_column_clears_the_field(senior, fields):
    account = AccountFactory(client_code="SH04", description="Old notes.")
    AccountCustomFieldValue.objects.create(
        account=account, definition=fields["text"], value="Dana"
    )

    apply_sheet(
        workbook(
            ["Client Code", "Client Name", "Description", "Go-live owner"],
            [["SH04", account.client_name, "", ""]],
        ),
        senior,
    )

    account.refresh_from_db()
    assert account.description == ""
    assert account.custom_field_values.get().value == ""


def test_an_update_missing_the_stage_column_is_not_told_it_becomes_active(senior):
    AccountFactory(client_code="SH05")
    batch = apply_sheet(
        workbook(["Client Code", "Client Name"], [["SH05", "Same Co"]]), senior
    )

    warnings = batch.rows.get().warnings
    assert not any(w["column"] == "Lifecycle stage" for w in warnings)


# ---------------------------------------------------------------------------
# Values
# ---------------------------------------------------------------------------


def people():
    return {"by_email": {}, "by_name": {}}


def test_a_link_without_a_scheme_is_given_https(db):
    rows, _ = parse(
        io.BytesIO(workbook(["Client Code", "Client Name", "Tenant URL"],
                            [["SH06", "Co", "acme.tenant.example"]])),
        people=people(), filename="s.xlsx",
    )
    assert rows[0].errors == []
    assert rows[0].cleaned["tenant_url"] == "https://acme.tenant.example"


def test_something_that_is_not_a_link_is_an_error_on_the_row(db):
    rows, _ = parse(
        io.BytesIO(workbook(["Client Code", "Client Name", "Drive URL"],
                            [["SH07", "Co", "not a link at all"]])),
        people=people(), filename="s.xlsx",
    )
    assert [e["column"] for e in rows[0].errors] == ["Drive URL"]


def test_a_dropdown_value_outside_its_options_is_an_error(fields):
    rows, _ = parse(
        io.BytesIO(workbook(["Client Code", "Client Name", "Segment"],
                            [["SH08", "Co", "Telecom"]])),
        people=people(), filename="s.xlsx", custom_fields=active(fields),
    )
    assert rows[0].errors == [
        {"column": "Segment", "message": "'Telecom' is not one of: Banking, Government, Retail."}
    ]


def test_a_dropdown_value_matches_whatever_its_case(fields):
    rows, _ = parse(
        io.BytesIO(workbook(["Client Code", "Client Name", "segment"],
                            [["SH09", "Co", "government"]])),
        people=people(), filename="s.xlsx", custom_fields=active(fields),
    )
    assert rows[0].cleaned["custom_fields"] == {str(fields["choice"].pk): "Government"}


def test_free_text_starting_with_equals_exports_as_text(senior):
    AccountFactory(client_code="SH10", description="=== renewal notes")
    content = build_accounts_xlsx(Account.objects.all(), include_health=False)
    sheet = load_workbook(io.BytesIO(content)).worksheets[0]

    cell = sheet.cell(row=2, column=len(COLUMNS))
    assert cell.value == "=== renewal notes"
    assert cell.data_type == "s"


# ---------------------------------------------------------------------------
# Naming a custom field
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("name", ["Tenant URL", "notes", "LinkedIn"])
def test_a_custom_field_may_not_take_a_built_in_column_name(db, name):
    form = CustomFieldDefinitionForm(data={"name": name, "field_type": "TEXT"})
    assert not form.is_valid()
    assert "import sheet" in form.errors["name"][0]


def test_a_custom_field_may_not_differ_from_another_only_by_punctuation(fields):
    form = CustomFieldDefinitionForm(data={"name": "Go live owner", "field_type": "TEXT"})
    assert not form.is_valid()
    assert "Go-live owner" in form.errors["name"][0]

"""The License edition.

License is the dPhish product the client is on - Advantage, Complete, Drill,
Discover, Essential or Phishing - and not a seat count. It briefly was the
latter, which earns a test rather than a comment: the two live side by side in
the same sheet, and confusing them again would put a word in an integer column
and lose the edition entirely.
"""

from __future__ import annotations

import io

import pytest
from django.urls import reverse

from csm.accounts.models import Account
from csm.core.enums import LicenseType
from csm.importer.parsing import COLUMNS, parse
from tests.factories import AccountFactory, UserFactory
from tests.test_importer import a_row, sheet
from tests.test_views import PASSWORD, login


def a_licensed_row(edition, **overrides):
    """The shared row with its License cell set. HEADERS already has the column;
    appending a second one would leave the parser reading the first."""
    return a_row(License=edition, **overrides)


@pytest.fixture
def senior(db):
    user = UserFactory(senior_csm=True, email="lic-senior@dphish.com")
    user.set_password(PASSWORD)
    user.save(update_fields=["password"])
    return user


@pytest.fixture
def emails(db):
    """The importer's people index, with the users a_row() names."""
    from csm.importer.services import _people

    for email in ("imp-csm@dphish.com", "imp-acm@dphish.com"):
        UserFactory(email=email, csm=True)
    return _people()


# ---------------------------------------------------------------------------
# The six editions
# ---------------------------------------------------------------------------


def test_the_six_editions_are_the_whole_set():
    assert [label for _, label in LicenseType.choices] == [
        "Advantage", "Complete", "Drill", "Discover", "Essential", "Phishing",
    ]


def test_license_and_the_seat_count_are_different_columns():
    """Both come from the same sheet, so neither may swallow the other."""
    assert "license" in COLUMNS["license_type"]
    assert "license" not in COLUMNS["boq_users_count"]
    assert "boq users count" in COLUMNS["boq_users_count"]


# ---------------------------------------------------------------------------
# Importing
# ---------------------------------------------------------------------------


def test_an_edition_imports_by_its_label(emails):
    rows, unknown = parse(
        sheet([a_licensed_row("Complete")]), people=emails
    )

    assert unknown == []
    assert rows[0].errors == []
    assert rows[0].cleaned["license_type"] == LicenseType.COMPLETE


def test_the_match_ignores_case(emails):
    rows, _ = parse(sheet([a_licensed_row("essential")]), people=emails)
    assert rows[0].cleaned["license_type"] == LicenseType.ESSENTIAL


def test_an_unknown_edition_is_an_error_listing_the_real_ones(emails):
    rows, _ = parse(sheet([a_licensed_row("Platinum")]), people=emails)

    message = next(e["message"] for e in rows[0].errors if e["column"] == "License")
    assert "Platinum" in message
    assert "Advantage" in message and "Phishing" in message


def test_a_blank_edition_is_allowed(emails):
    """Not every account has one recorded, and that is a fact, not an error."""
    rows, _ = parse(sheet([a_licensed_row("")]), people=emails)

    assert rows[0].errors == []
    assert rows[0].cleaned["license_type"] == ""


def test_a_number_in_the_license_column_no_longer_means_seats(emails):
    """It used to. Putting 500 there now names no edition, so it fails."""
    rows, _ = parse(sheet([a_licensed_row("500")]), people=emails)
    assert any(e["column"] == "License" for e in rows[0].errors)


def test_the_seat_count_still_imports_separately(emails):
    rows, _ = parse(
        sheet([a_licensed_row("Drill", **{"BOQ Users Count": 750})]),
        people=emails,
    )

    assert rows[0].errors == []
    assert rows[0].cleaned["license_type"] == LicenseType.DRILL
    assert rows[0].cleaned["boq_users_count"] == 750


# ---------------------------------------------------------------------------
# The templates
# ---------------------------------------------------------------------------


def test_the_template_offers_license_as_its_own_column():
    from csm.importer.parsing import read_table
    from csm.importer.template import build_template, header_label

    assert header_label("license_type") == "License"
    assert header_label("boq_users_count") == "BOQ Users Count"

    header_map, body, unknown = read_table(
        io.BytesIO(build_template()), filename="t.xlsx"
    )
    assert unknown == []
    assert "license_type" in header_map
    assert "boq_users_count" in header_map


def test_the_template_instructions_list_the_editions():
    from openpyxl import load_workbook

    from csm.importer.template import build_template

    workbook = load_workbook(io.BytesIO(build_template()))
    notes = " ".join(
        str(cell.value)
        for row in workbook["Instructions"].iter_rows()
        for cell in row
        if cell.value
    )
    for label in ("Advantage", "Complete", "Drill", "Discover", "Essential", "Phishing"):
        assert label in notes


# ---------------------------------------------------------------------------
# The interface
# ---------------------------------------------------------------------------


def test_the_dashboard_has_a_license_card(client, senior):
    AccountFactory(client_code="L1", csm=senior, license_type=LicenseType.COMPLETE)
    AccountFactory(client_code="L2", csm=senior, license_type=LicenseType.COMPLETE)
    AccountFactory(client_code="L3", csm=senior, license_type=LicenseType.DRILL)
    login(client, senior)

    response = client.get(reverse("dashboard:analytics"))
    card = {c.id: c.data for c in response.context["cards"]}["license_distribution"]
    counts = dict(card["rows"])

    assert counts["Complete"] == 2
    assert counts["Drill"] == 1
    assert counts["Phishing"] == 0, "an edition nobody is on still gets a bar"
    assert "License editions" in response.content.decode()


def test_accounts_with_no_edition_get_their_own_bar(client, senior):
    """Otherwise the chart adds up to fewer accounts than the total beside it."""
    AccountFactory(client_code="L4", csm=senior, license_type=LicenseType.ESSENTIAL)
    AccountFactory(client_code="L5", csm=senior, license_type="")
    login(client, senior)

    response = client.get(reverse("dashboard:analytics"))
    card = {c.id: c.data for c in response.context["cards"]}["license_distribution"]
    counts = dict(card["rows"])

    assert counts["Not set"] == 1
    assert sum(counts.values()) == 2


def test_the_account_page_shows_the_edition(client, senior):
    account = AccountFactory(
        client_code="L6", client_name="Acme", csm=senior,
        license_type=LicenseType.ADVANTAGE,
    )
    login(client, senior)

    body = client.get(reverse("accounts:detail", args=[account.pk])).content.decode()
    assert "Advantage" in body


def test_the_account_list_can_be_filtered_by_edition(client, senior):
    AccountFactory(
        client_code="L7", client_name="Complete Co", csm=senior,
        license_type=LicenseType.COMPLETE,
    )
    AccountFactory(
        client_code="L8", client_name="Drill Co", csm=senior,
        license_type=LicenseType.DRILL,
    )
    login(client, senior)

    body = client.get(
        reverse("accounts:list"), {"license": LicenseType.COMPLETE}
    ).content.decode()

    assert "Complete Co" in body
    assert "Drill Co" not in body


def test_the_edition_is_editable_and_stored(client, senior):
    account = AccountFactory(client_code="L9", csm=senior)
    login(client, senior)

    client.post(
        reverse("accounts:edit", args=[account.pk]),
        {
            "client_code": account.client_code, "client_name": account.client_name,
            "license_type": LicenseType.PHISHING, "deployment": "CLOUD",
            "account_class": "B", "size": "MID", "lifecycle_stage": "ACTIVE",
            "boq_users_count": 10, "added_users_count": 0, "renewal_count": 0,
            "events_count": 0, "digital_content_count": 0,
            "g2_review_status": "NONE", "gartner_review_status": "NONE",
        },
    )

    assert Account.objects.get(pk=account.pk).license_type == LicenseType.PHISHING

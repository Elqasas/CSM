"""CSV uploads and the contacts import.

Two things earn most of these tests. CSV has more ways to arrive malformed than
xlsx does - byte-order marks, a semicolon delimiter from a European Excel, a
non-UTF-8 encoding - and each of them fails by parsing as *something*, not by
raising. And a contacts file must never create an account: doing so would
produce a client with a name and nothing else that still counts on every
dashboard card.
"""

from __future__ import annotations

import datetime as dt
import io

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.urls import reverse

from csm.accounts.models import Account
from csm.contacts.models import Contact
from csm.core.enums import CommunicationChannel, Level
from csm.importer import contact_services, services
from csm.importer.contact_parsing import CONTACT_COLUMNS, parse_contacts
from csm.importer.csv_template import build_accounts_csv, build_contacts_csv
from csm.importer.models import ImportBatch, ImportKind, ImportStatus
from csm.importer.parsing import parse, read_table
from csm.importer.template import build_template
from tests.factories import AccountFactory, UserFactory
from tests.test_importer import HEADERS, a_row
from tests.test_views import PASSWORD, login

CONTACT_HEADERS = [
    "Client Code", "Full Name", "Title", "Email", "Phone",
    "Influence", "Interest", "Channels", "Preferred Channel", "Focal Rank", "Notes",
]


def csv_bytes(headers, rows, *, delimiter=",", bom=False, encoding="utf-8") -> io.BytesIO:
    """Write real CSV, quoting included.

    Joining on the delimiter by hand is not the same thing: a cell such as
    "Email, Teams" then splits into two columns and every field after it shifts
    left, which surfaces as unrelated validation errors several columns away.
    """
    import csv

    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=delimiter, lineterminator=chr(13)+chr(10))
    writer.writerow(headers)
    for row in rows:
        writer.writerow(["" if c is None else c for c in row])

    text = buffer.getvalue()
    if bom:
        text = "﻿" + text
    return io.BytesIO(text.encode(encoding))


def a_contact(**overrides) -> list:
    values = {
        "Client Code": "ACME01", "Full Name": "Dana Farouk", "Title": "CISO",
        "Email": "dana@acme.example", "Phone": "+20 100", "Influence": "High",
        "Interest": "Medium", "Channels": "Email, Teams",
        "Preferred Channel": "Email", "Focal Rank": "1", "Notes": "",
    }
    values.update(overrides)
    return [values[h] for h in CONTACT_HEADERS]


@pytest.fixture
def staff(db):
    # The same addresses tests.test_importer.a_row() puts in the CSM/ACM
    # columns, so a row reused from there resolves rather than erroring.
    users = {
        "senior": UserFactory(senior_csm=True, email="csv-senior@dphish.com"),
        "csm": UserFactory(csm=True, email="imp-csm@dphish.com"),
        "acm": UserFactory(acm=True, email="imp-acm@dphish.com"),
    }
    for user in users.values():
        user.set_password(PASSWORD)
        user.save(update_fields=["password"])
    return users


@pytest.fixture
def emails(staff):
    from csm.importer.services import _people

    return _people()


@pytest.fixture
def account(db):
    return AccountFactory(client_code="ACME01", client_name="Acme Industrial")


def make_batch(staff, content: bytes, name: str, kind=ImportKind.ACCOUNTS) -> ImportBatch:
    batch = ImportBatch.objects.create(
        upload=SimpleUploadedFile(name, content),
        original_name=name,
        kind=kind,
        uploaded_by=staff["senior"],
        baseline_date=dt.date(2026, 9, 21),
    )
    module = contact_services if kind == ImportKind.CONTACTS else services
    return module.validate(batch)


# ---------------------------------------------------------------------------
# CSV reading
# ---------------------------------------------------------------------------


def test_a_plain_csv_parses_like_a_spreadsheet(emails):
    rows, unknown = parse(
        csv_bytes(HEADERS, [a_row()]), people=emails, filename="portfolio.csv"
    )
    assert unknown == []
    assert rows[0].errors == []
    assert rows[0].cleaned["country"] == "EG"


def test_a_byte_order_mark_does_not_break_the_first_header(emails):
    """Excel writes one. Without stripping it, 'Client Code' matches nothing and
    the sheet is refused for having no required columns."""
    rows, _ = parse(
        csv_bytes(HEADERS, [a_row()], bom=True),
        people=emails, filename="portfolio.csv",
    )
    assert rows[0].client_code == "ACME01"
    assert rows[0].errors == []


def test_a_semicolon_delimited_csv_is_read_correctly(emails):
    """Excel emits these in locales whose decimal separator is a comma. Read as
    comma-separated it becomes one enormous column and validates as empty."""
    rows, _ = parse(
        csv_bytes(HEADERS, [a_row(**{"Partner": "Regional Reseller"})], delimiter=";"),
        people=emails, filename="portfolio.csv",
    )
    assert len(rows) == 1
    assert rows[0].client_code == "ACME01"
    assert rows[0].errors == []


def test_a_non_utf8_csv_is_decoded_rather_than_refused(emails):
    rows, _ = parse(
        csv_bytes(HEADERS, [a_row(**{"Location": "Zürich"})], encoding="cp1252"),
        people=emails, filename="portfolio.csv",
    )
    assert rows[0].cleaned["location"] == "Zürich"


def test_an_empty_csv_is_refused_with_a_readable_message(emails):
    from csm.importer.parsing import WorkbookError

    with pytest.raises(WorkbookError, match="empty"):
        parse(io.BytesIO(b"   "), people=emails, filename="portfolio.csv")


def test_comment_lines_are_not_data(emails):
    body = b"# a note about this file\r\n" + csv_bytes(HEADERS, [a_row()]).getvalue()
    rows, _ = parse(io.BytesIO(body), people=emails, filename="portfolio.csv")
    assert len(rows) == 1


# ---------------------------------------------------------------------------
# The CSV templates
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "builder, columns",
    [(build_accounts_csv, None), (build_contacts_csv, CONTACT_COLUMNS)],
)
def test_a_csv_template_round_trips_through_its_own_parser(builder, columns):
    """A template that drifts from the parser produces a file that validates as
    empty, which is the failure nobody notices."""
    kwargs = {"columns": columns} if columns else {}
    header_map, body, unknown = read_table(
        io.BytesIO(builder()), filename="template.csv", **kwargs,
        **({"required": ("client_code", "full_name")} if columns else {}),
    )

    assert unknown == [], f"template emits headers the parser ignores: {unknown}"
    assert body == [], "the example row must be commented out, not importable"


def test_the_accounts_csv_template_imports_nothing_by_itself(emails):
    rows, _ = parse(
        io.BytesIO(build_accounts_csv()), people=emails, filename="t.csv"
    )
    assert rows == []


def test_the_contacts_csv_template_imports_nothing_by_itself(db):
    rows, _ = parse_contacts(
        io.BytesIO(build_contacts_csv()), known_codes=set(), filename="t.csv"
    )
    assert rows == []


def test_the_contacts_xlsx_template_uses_the_contact_columns():
    from openpyxl import load_workbook

    from csm.importer.contact_parsing import CONTACT_REQUIRED_COLUMNS

    content = build_template(CONTACT_COLUMNS, CONTACT_REQUIRED_COLUMNS)
    ws = load_workbook(io.BytesIO(content)).worksheets[0]
    headers = [str(c.value) for c in ws[1]]

    assert any("Full Name" in h for h in headers)
    assert not any("Lifecycle" in h for h in headers)


# ---------------------------------------------------------------------------
# Contacts parsing
# ---------------------------------------------------------------------------


def test_a_clean_contact_row_parses(db, account):
    rows, unknown = parse_contacts(
        csv_bytes(CONTACT_HEADERS, [a_contact()]),
        known_codes={"ACME01"}, filename="contacts.csv",
    )
    assert unknown == []
    assert rows[0].errors == []
    assert rows[0].cleaned["influence"] == Level.HIGH
    assert rows[0].cleaned["channels"] == [
        CommunicationChannel.EMAIL, CommunicationChannel.TEAMS
    ]
    assert rows[0].cleaned["focal_rank"] == 1


def test_a_contact_for_an_unknown_account_is_an_error(db):
    rows, _ = parse_contacts(
        csv_bytes(CONTACT_HEADERS, [a_contact(**{"Client Code": "NOPE"})]),
        known_codes={"ACME01"}, filename="contacts.csv",
    )
    assert any("never creates an account" in e["message"] for e in rows[0].errors)


def test_a_blank_assessment_imports_as_unassessed_not_medium(db):
    rows, _ = parse_contacts(
        csv_bytes(CONTACT_HEADERS, [a_contact(**{"Influence": "", "Interest": ""})]),
        known_codes={"ACME01"}, filename="contacts.csv",
    )
    assert rows[0].errors == []
    assert rows[0].cleaned["influence"] is None
    assert rows[0].cleaned["interest"] is None
    assert any("triage tray" in w["message"] for w in rows[0].warnings)


def test_half_an_assessment_is_flagged_as_still_unplaced(db):
    rows, _ = parse_contacts(
        csv_bytes(CONTACT_HEADERS, [a_contact(**{"Interest": ""})]),
        known_codes={"ACME01"}, filename="contacts.csv",
    )
    assert any("no position on the grid" in w["message"] for w in rows[0].warnings)


def test_channels_split_on_whatever_separator_was_used(db):
    for text in ("Email, Teams", "Email;Teams", "Email / Teams", "Email|Teams"):
        rows, _ = parse_contacts(
            csv_bytes(CONTACT_HEADERS, [a_contact(**{"Channels": text})]),
            known_codes={"ACME01"}, filename="contacts.csv",
        )
        assert rows[0].cleaned["channels"] == [
            CommunicationChannel.EMAIL, CommunicationChannel.TEAMS
        ], text


def test_a_preferred_channel_outside_the_selected_ones_is_an_error(db):
    rows, _ = parse_contacts(
        csv_bytes(CONTACT_HEADERS, [a_contact(**{"Preferred Channel": "Phone"})]),
        known_codes={"ACME01"}, filename="contacts.csv",
    )
    assert any(e["column"] == "Preferred channel" for e in rows[0].errors)


def test_a_focal_point_without_a_rank_is_an_error(db):
    rows, _ = parse_contacts(
        csv_bytes(CONTACT_HEADERS, [a_contact(**{"Focal Rank": "yes"})]),
        known_codes={"ACME01"}, filename="contacts.csv",
    )
    assert any("rank of 1 or 2" in e["message"] for e in rows[0].errors)


def test_two_rows_claiming_the_same_focal_rank_is_caught_before_the_write(db):
    """The database keeps one contact per rank per account, so this would
    otherwise fail the whole import instead of naming the row."""
    rows, _ = parse_contacts(
        csv_bytes(CONTACT_HEADERS, [
            a_contact(),
            a_contact(**{"Full Name": "Omar Said"}),
        ]),
        known_codes={"ACME01"}, filename="contacts.csv",
    )
    assert rows[0].errors == []
    assert any("already claims focal rank 1" in e["message"] for e in rows[1].errors)


# ---------------------------------------------------------------------------
# Contacts apply
# ---------------------------------------------------------------------------


def test_applying_contacts_creates_them_against_the_right_account(db, staff, account):
    batch = make_batch(
        staff, csv_bytes(CONTACT_HEADERS, [a_contact()]).getvalue(),
        "contacts.csv", ImportKind.CONTACTS,
    )
    assert Contact.objects.count() == 0, "preview must not write"

    result = contact_services.apply(batch, actor=staff["senior"])
    contact = Contact.objects.get()

    assert result["created"] == 1
    assert contact.account == account
    assert contact.full_name == "Dana Farouk"
    assert contact.influence == Level.HIGH
    # High influence with Medium interest is Keep satisfied, not Manage closely:
    # Medium counts as the low half of the interest axis.
    assert contact.quadrant == "KEEP_SATISFIED"
    assert set(contact.channels.values_list("channel", flat=True)) == {
        CommunicationChannel.EMAIL, CommunicationChannel.TEAMS
    }


def test_an_unassessed_imported_contact_lands_in_the_triage_tray(db, staff, account):
    batch = make_batch(
        staff,
        csv_bytes(CONTACT_HEADERS, [a_contact(**{"Influence": "", "Interest": ""})]).getvalue(),
        "contacts.csv", ImportKind.CONTACTS,
    )
    contact_services.apply(batch, actor=staff["senior"])

    contact = Contact.objects.get()
    assert contact.quadrant is None
    assert Contact.objects.needs_triage().count() == 1


def test_reimporting_a_contact_updates_rather_than_duplicating(db, staff, account):
    contact_services.apply(
        make_batch(staff, csv_bytes(CONTACT_HEADERS, [a_contact()]).getvalue(),
                   "c.csv", ImportKind.CONTACTS),
        actor=staff["senior"],
    )
    batch = make_batch(
        staff,
        csv_bytes(CONTACT_HEADERS, [a_contact(**{"Title": "Group CISO"})]).getvalue(),
        "c.csv", ImportKind.CONTACTS,
    )
    assert batch.rows_to_update == 1

    result = contact_services.apply(batch, actor=staff["senior"])
    assert result["updated"] == 1
    assert Contact.objects.count() == 1
    assert Contact.objects.get().title == "Group CISO"


def test_a_contacts_file_never_creates_an_account(db, staff):
    batch = make_batch(
        staff,
        csv_bytes(CONTACT_HEADERS, [a_contact(**{"Client Code": "GHOST"})]).getvalue(),
        "contacts.csv", ImportKind.CONTACTS,
    )
    assert batch.rows_ready == 0
    assert Account.objects.count() == 0


# ---------------------------------------------------------------------------
# Through the views
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "kind, fmt", [("accounts", "csv"), ("accounts", "xlsx"),
                  ("contacts", "csv"), ("contacts", "xlsx")],
)
def test_all_four_templates_download(client, staff, kind, fmt):
    login(client, staff["senior"])
    response = client.get(reverse("importer:template_for", args=[kind, fmt]))

    assert response.status_code == 200
    assert f"{kind}-template.{fmt}" in response["Content-Disposition"]
    assert response.content


def test_the_contacts_wizard_runs_end_to_end(client, staff, account):
    login(client, staff["senior"])
    upload = SimpleUploadedFile(
        "contacts.csv", csv_bytes(CONTACT_HEADERS, [a_contact()]).getvalue(),
        content_type="text/csv",
    )

    client.post(
        reverse("importer:upload"), {"upload": upload, "kind": "CONTACTS"}, follow=True
    )
    assert Contact.objects.count() == 0, "upload previews, it does not write"

    batch = ImportBatch.objects.get()
    assert batch.kind == ImportKind.CONTACTS

    client.post(reverse("importer:apply", args=[batch.pk]))
    batch.refresh_from_db()

    assert batch.status == ImportStatus.APPLIED
    assert Contact.objects.filter(full_name="Dana Farouk").exists()
    assert batch.rows.first().contact is not None


def test_an_accounts_file_uploaded_as_contacts_fails_loudly(client, staff):
    """The kind is asked rather than sniffed, so a mismatch has to say so."""
    login(client, staff["senior"])
    upload = SimpleUploadedFile(
        "portfolio.csv", csv_bytes(HEADERS, [a_row()]).getvalue(), content_type="text/csv"
    )
    client.post(reverse("importer:upload"), {"upload": upload, "kind": "CONTACTS"})

    batch = ImportBatch.objects.get()
    assert batch.status == ImportStatus.FAILED
    assert "Full Name" in batch.parse_error


def test_a_csv_uploaded_as_a_batch_reaches_the_csv_reader(db, staff):
    """The batch path picks its reader from the filename.

    Parsing directly with an explicit filename does not exercise that: without
    it every upload went down the xlsx path and a CSV failed as "not a zip
    file", which is how this got past the unit tests the first time.
    """
    batch = make_batch(
        staff, csv_bytes(HEADERS, [a_row()]).getvalue(), "portfolio.csv"
    )

    assert batch.status == ImportStatus.VALIDATED, batch.parse_error
    assert batch.rows_ready == 1


def test_a_contacts_csv_uploaded_as_a_batch_reaches_the_csv_reader(db, staff, account):
    batch = make_batch(
        staff, csv_bytes(CONTACT_HEADERS, [a_contact()]).getvalue(),
        "contacts.csv", ImportKind.CONTACTS,
    )

    assert batch.status == ImportStatus.VALIDATED, batch.parse_error
    assert batch.rows_ready == 1

"""The Reports page: who may run what, what each report holds, and the period.

Builders are asserted directly through a ``ReportContext`` - the sheets are
plain lists, so no workbook has to be parsed to check a row - and the views are
asserted for access, scope and the file that comes back.
"""

from __future__ import annotations

import datetime as dt
import io
from decimal import Decimal

import pytest
import time_machine
from django.urls import reverse
from openpyxl import load_workbook

from csm.core.access.abilities import Ability
from csm.core.enums import (
    ActivityType,
    HealthStatus,
    ScoreTrigger,
    TaskStatus,
    TaskType,
)
from csm.health.models import HealthScoreHistory, HealthStatusOverride
from csm.reports import builders, periods, registry
from csm.reports.builders import ReportContext
from csm.tasks.generators import project_auto_tasks, renewal_key
from csm.tasks.models import Task
from csm.tasks.services import create_task
from tests.factories import AccountFactory, ActivityFactory, UserFactory
from tests.test_views import PASSWORD, login

TODAY = dt.date(2026, 10, 18)
UTC = dt.UTC


# ---------------------------------------------------------------------------
# Periods - pure
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("key", "today", "start", "end"),
    [
        ("this_month", dt.date(2026, 2, 10), dt.date(2026, 2, 1), dt.date(2026, 2, 28)),
        ("last_month", dt.date(2026, 1, 5), dt.date(2025, 12, 1), dt.date(2025, 12, 31)),
        ("this_quarter", dt.date(2026, 10, 18), dt.date(2026, 10, 1), dt.date(2026, 12, 31)),
        ("last_quarter", dt.date(2026, 2, 1), dt.date(2025, 10, 1), dt.date(2025, 12, 31)),
        ("next_quarter", dt.date(2026, 11, 30), dt.date(2027, 1, 1), dt.date(2027, 3, 31)),
        ("this_year", dt.date(2026, 6, 1), dt.date(2026, 1, 1), dt.date(2026, 12, 31)),
        ("last_year", dt.date(2026, 6, 1), dt.date(2025, 1, 1), dt.date(2025, 12, 31)),
        ("last_30", dt.date(2026, 3, 30), dt.date(2026, 3, 1), dt.date(2026, 3, 30)),
        ("next_30", dt.date(2026, 3, 1), dt.date(2026, 3, 1), dt.date(2026, 3, 30)),
        ("next_90", dt.date(2026, 1, 1), dt.date(2026, 1, 1), dt.date(2026, 3, 31)),
    ],
)
def test_every_preset_across_its_boundary(key, today, start, end):
    period = periods.resolve(key, "", "", today=today)
    assert (period.start, period.end) == (start, end)


def test_a_preset_ignores_the_date_boxes():
    period = periods.resolve("this_month", "2020-01-01", "2020-01-02", today=TODAY)
    assert period.start == dt.date(2026, 10, 1)


def test_a_custom_range_reads_both_boxes():
    period = periods.resolve(periods.CUSTOM, "2026-07-01", "2026-09-30", today=TODAY)
    assert (period.start, period.end, period.days) == (
        dt.date(2026, 7, 1),
        dt.date(2026, 9, 30),
        92,
    )


@pytest.mark.parametrize(
    ("key", "start", "end", "message"),
    [
        (periods.CUSTOM, "2026-09-30", "2026-07-01", "before the start"),
        (periods.CUSTOM, "", "2026-07-01", "start date"),
        (periods.CUSTOM, "2026-07-01", "not a date", "end date"),
        (periods.CUSTOM, "2023-01-01", "2026-01-01", "two years"),
        ("fortnight", "", "", "Choose a period"),
    ],
)
def test_a_range_that_cannot_be_built_says_why(key, start, end, message):
    with pytest.raises(periods.PeriodError, match=message):
        periods.resolve(key, start, end, today=TODAY)


def test_two_years_exactly_is_allowed():
    start = dt.date(2025, 1, 1)
    end = start + dt.timedelta(days=periods.MAX_DAYS - 1)
    assert periods.resolve(periods.CUSTOM, str(start), str(end), today=TODAY).end == end


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


@pytest.fixture
def staff(db):
    users = {
        "csm": UserFactory(csm=True, email="rep-csm@dphish.com", full_name="Rana CSM"),
        "other": UserFactory(csm=True, email="rep-other@dphish.com", full_name="Omar CSM"),
        "senior": UserFactory(senior_csm=True, email="rep-senior@dphish.com"),
        "acm": UserFactory(acm=True, email="rep-acm@dphish.com"),
        "manager": UserFactory(manager=True, email="rep-manager@dphish.com"),
        "admin": UserFactory(admin=True, email="rep-admin@dphish.com"),
    }
    for user in users.values():
        user.set_password(PASSWORD)
        user.save(update_fields=["password"])
    return users


@pytest.fixture
def book(staff):
    return {
        "mine": AccountFactory(
            client_code="REP01", client_name="Mine Co", csm=staff["csm"], acm=staff["acm"]
        ),
        "theirs": AccountFactory(client_code="REP02", client_name="Theirs Co", csm=staff["other"]),
    }


def context(user, start, end, today=TODAY):
    period = periods.Period(start=start, end=end, label="test")
    return ReportContext(
        ability=Ability(user),
        period=period,
        today=today,
        generated_at=dt.datetime(2026, 10, 18, 9, 0, tzinfo=UTC),
    )


def sheet(sheets, title):
    return next(s for s in sheets if s.title == title)


def column(found, header):
    index = found.headers.index(header)
    return [row[index] for row in found.rows]


def done(account, title, due, completed_at):
    task = create_task(
        account=account, task_type=TaskType.MANUAL_NEXT_ACTION, title=title, due_date=due
    )
    Task.objects.filter(pk=task.pk).update(status=TaskStatus.DONE, completed_at=completed_at)
    return task


# ---------------------------------------------------------------------------
# Access
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("who", ["csm", "senior", "manager", "admin"])
def test_every_report_is_offered_to_the_csm_side(staff, who):
    keys = {r.key for r in registry.visible_reports(Ability(staff[who]))}
    assert keys == {r.key for r in registry.REPORTS}


def test_an_acm_is_offered_the_two_data_sheets_and_nothing_else(staff):
    keys = {r.key for r in registry.visible_reports(Ability(staff["acm"]))}
    assert keys == {"accounts", "contacts"}


def test_an_acm_cannot_download_a_report_by_url(client, staff, book):
    login(client, staff["acm"])
    for report in registry.REPORTS:
        response = client.get(reverse("reports:download", args=[report.key]))
        expected = 200 if report.key in ("accounts", "contacts") else 403
        assert response.status_code == expected, report.key


def test_an_acms_page_shows_no_report_card(client, staff):
    login(client, staff["acm"])
    page = client.get(reverse("reports:index")).content.decode()
    assert reverse("reports:download", args=["accounts"]) in page
    assert reverse("reports:download", args=["health"]) not in page
    assert reverse("reports:download", args=["work"]) not in page


def test_reports_are_reached_from_settings_not_the_top_menu(client, staff):
    login(client, staff["csm"])
    assert reverse("reports:index") in client.get(reverse("settings")).content.decode()
    page = client.get(reverse("dashboard:index")).content.decode()
    assert reverse("reports:index") not in page


def test_the_page_opens_on_the_tab_it_is_asked_for(client, staff):
    login(client, staff["csm"])
    page = client.get(reverse("reports:index"), {"report": "health"}).content.decode()
    assert 'class="tab-pane fade show active"\n           id="pane-health"' in page
    # A report key inside a shared tab opens that tab.
    page = client.get(reverse("reports:index"), {"report": "contacts"}).content.decode()
    assert 'class="tab-pane fade show active"\n           id="pane-data"' in page


def test_an_unknown_report_is_a_404(client, staff):
    login(client, staff["csm"])
    assert client.get(reverse("reports:download", args=["payroll"])).status_code == 404


def test_the_old_export_address_redirects_to_reports(client, staff):
    login(client, staff["csm"])
    response = client.get(reverse("importer:export"))
    assert response.status_code == 302
    assert response["Location"] == reverse("reports:index")


# ---------------------------------------------------------------------------
# Downloads
# ---------------------------------------------------------------------------


def test_a_report_downloads_with_an_about_sheet(client, staff, book):
    login(client, staff["csm"])
    response = client.get(
        reverse("reports:download", args=["work"]),
        {"period": "custom", "from": "2026-09-01", "to": "2026-09-30", "show": "completed"},
    )
    assert response.status_code == 200
    assert 'filename="work_2026-09-01_to_2026-09-30.xlsx"' in response[
        "Content-Disposition"
    ]
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.sheetnames == ["Summary by CSM", "Tasks and activities", "About"]
    about = {row[0]: row[1] for row in workbook["About"].iter_rows(values_only=True)}
    assert about["Report"] == "Tasks & activities"
    assert about["Show"] == "Completed only"
    assert about["Accounts in scope"] == 1


def test_a_bad_range_reopens_the_page_with_the_error_on_its_card(client, staff):
    login(client, staff["csm"])
    response = client.get(
        reverse("reports:download", args=["work"]),
        {"period": "custom", "from": "2026-09-30", "to": "2026-09-01"},
    )
    assert response.status_code == 400
    page = response.content.decode()
    assert "The end date is before the start date." in page
    assert 'value="2026-09-30"' in page


def test_every_report_builds_for_every_role_that_may_run_it(client, staff, book):
    for who in ("csm", "senior", "manager", "admin", "acm"):
        client.logout()
        login(client, staff[who])
        for report in registry.visible_reports(Ability(staff[who])):
            response = client.get(reverse("reports:download", args=[report.key]))
            assert response.status_code == 200, (who, report.key)
            load_workbook(io.BytesIO(response.content))


# ---------------------------------------------------------------------------
# Tasks and activities
# ---------------------------------------------------------------------------

SEPT = (dt.date(2026, 9, 1), dt.date(2026, 9, 30))


def work(user, start, end, show="all", today=TODAY):
    ctx = context(user, start, end, today=today)
    ctx.options["show"] = show
    sheets = builders.work(ctx)
    return sheet(sheets, "Tasks and activities"), sheet(sheets, "Summary by CSM")


def test_a_completed_task_is_one_row_its_activity_record(staff, book):
    task = done(
        book["mine"], "QBR", dt.date(2026, 9, 10), dt.datetime(2026, 9, 14, 8, tzinfo=UTC)
    )
    ActivityFactory(
        account=book["mine"], task=task, occurred_on=dt.date(2026, 9, 14),
        activity_type=ActivityType.MEETING, created_by=staff["csm"],
    )

    rows, _ = work(staff["csm"], *SEPT)
    assert column(rows, "Title") == ["QBR"]
    assert column(rows, "Kind") == ["Meeting"]
    assert column(rows, "Due date") == [dt.date(2026, 9, 10)]
    assert column(rows, "Days late") == [4]


def test_completed_rows_cover_records_without_a_task_and_tasks_without_a_record(staff, book):
    ActivityFactory(
        account=book["mine"], occurred_on=dt.date(2026, 9, 3), created_by=staff["csm"],
        activity_type=ActivityType.MEETING, subject="Teams sync", source="TEAMS",
    )
    done(book["mine"], "Old task", dt.date(2026, 9, 10), dt.datetime(2026, 9, 9, 8, tzinfo=UTC))

    rows, _ = work(staff["csm"], *SEPT)
    assert column(rows, "Title") == ["Teams sync", "Old task"]
    assert column(rows, "Source") == ["Teams", "Platform"]
    assert column(rows, "Days late") == [None, 0]


def test_voided_records_and_other_portfolios_are_left_out(staff, book):
    ActivityFactory(
        account=book["mine"], occurred_on=dt.date(2026, 9, 5), created_by=staff["csm"],
        subject="Voided", is_voided=True,
    )
    ActivityFactory(account=book["theirs"], occurred_on=dt.date(2026, 9, 3), created_by=staff["other"])
    done(book["theirs"], "Theirs", dt.date(2026, 9, 10), dt.datetime(2026, 9, 9, 8, tzinfo=UTC))

    rows, _ = work(staff["csm"], *SEPT)
    assert rows.rows == []


def test_a_note_is_not_client_contact(staff, book):
    ActivityFactory(
        account=book["mine"], occurred_on=dt.date(2026, 9, 4), created_by=staff["csm"],
        activity_type=ActivityType.NOTE, subject="Internal",
    )
    rows, summary = work(staff["csm"], *SEPT)
    assert column(rows, "Client contact") == ["No"]
    assert column(summary, "Client contacts") == [0]


def test_the_period_ends_at_midnight_in_cairo_not_utc(staff, book):
    # 22:30 UTC on 30 September is 01:30 on 1 October in Cairo (UTC+3).
    done(book["mine"], "October", dt.date(2026, 10, 1), dt.datetime(2026, 9, 30, 22, 30, tzinfo=UTC))
    # 21:30 UTC on 31 August is 00:30 on 1 September in Cairo.
    done(book["mine"], "September", dt.date(2026, 9, 1), dt.datetime(2026, 8, 31, 21, 30, tzinfo=UTC))

    rows, _ = work(staff["csm"], *SEPT, show="completed")
    assert column(rows, "Title") == ["September"]


def test_a_cancelled_task_is_neither_done_nor_upcoming(staff, book):
    task = create_task(
        account=book["mine"],
        task_type=TaskType.MANUAL_NEXT_ACTION,
        title="Dropped",
        due_date=dt.date(2026, 9, 10),
    )
    Task.objects.filter(pk=task.pk).update(status=TaskStatus.CANCELLED)

    rows, _ = work(staff["csm"], *SEPT)
    assert rows.rows == []


def test_upcoming_lists_open_tasks_and_projects_the_renewal_reminders(staff, book):
    account = book["mine"]
    account.contract_end_date = dt.date(2027, 1, 31)
    account.save()
    create_task(
        account=account,
        task_type=TaskType.MANUAL_NEXT_ACTION,
        title="Send the deck",
        due_date=dt.date(2026, 11, 5),
    )

    rows, summary = work(staff["csm"], TODAY, dt.date(2026, 12, 31), show="upcoming")
    # The 30 day reminder falls on 1 January, after the period.
    assert column(rows, "Title") == [
        "Renewal in 90 days: Mine Co",
        "Send the deck",
        "Renewal in 60 days: Mine Co",
    ]
    assert column(rows, "Status") == ["Projected", "Open", "Projected"]
    assert summary.headers == ["CSM", "Overdue", "Open", "Projected"]
    assert summary.rows == [["Rana CSM", 0, 1, 2]]


def test_a_past_period_shows_what_is_still_open_as_overdue(staff, book):
    create_task(
        account=book["mine"],
        task_type=TaskType.MANUAL_NEXT_ACTION,
        title="Slipped",
        due_date=dt.date(2026, 9, 10),
    )
    rows, _ = work(staff["csm"], *SEPT, show="upcoming")
    assert column(rows, "Status") == ["Overdue"]


def test_the_show_option_picks_the_half(staff, book):
    done(book["mine"], "Done", dt.date(2026, 10, 2), dt.datetime(2026, 10, 2, 8, tzinfo=UTC))
    create_task(
        account=book["mine"],
        task_type=TaskType.MANUAL_NEXT_ACTION,
        title="Planned",
        due_date=dt.date(2026, 10, 25),
    )
    october = (dt.date(2026, 10, 1), dt.date(2026, 10, 31))

    assert column(work(staff["csm"], *october, show="all")[0], "Title") == ["Done", "Planned"]
    assert column(work(staff["csm"], *october, show="completed")[0], "Title") == ["Done"]
    assert column(work(staff["csm"], *october, show="upcoming")[0], "Title") == ["Planned"]


def test_an_unknown_show_value_is_refused_not_guessed(client, staff):
    login(client, staff["csm"])
    response = client.get(reverse("reports:download", args=["work"]), {"show": "everything"})
    assert response.status_code == 400
    assert "Choose what to show" in response.content.decode()


def test_the_projection_skips_a_reminder_that_already_exists(staff, book):
    account = book["mine"]
    account.contract_end_date = dt.date(2027, 1, 31)
    account.save()
    create_task(
        account=account,
        task_type=TaskType.AUTO_RENEWAL,
        title="Created early by hand",
        due_date=dt.date(2026, 11, 2),
        auto_key=renewal_key(account.pk, account.contract_end_date, 90),
    )

    projected = project_auto_tasks(
        Ability(staff["csm"]).accounts(), TODAY, dt.date(2026, 12, 31), as_of=TODAY
    )
    assert [p.due_date for p in projected] == [dt.date(2026, 12, 2)]


def test_the_projection_never_reaches_today_or_before(staff, book):
    account = book["mine"]
    account.contract_end_date = TODAY + dt.timedelta(days=30)
    account.save()

    projected = project_auto_tasks(
        Ability(staff["csm"]).accounts(), dt.date(2026, 1, 1), TODAY, as_of=TODAY
    )
    assert projected == []


def test_the_projection_skips_a_churned_account(staff):
    AccountFactory(csm=staff["csm"], churned=True, contract_end_date=dt.date(2027, 1, 31))
    projected = project_auto_tasks(
        Ability(staff["csm"]).accounts(), TODAY, dt.date(2027, 1, 31), as_of=TODAY
    )
    assert projected == []


def test_the_projection_includes_the_quarterly_meeting_lead(staff, book):
    account = book["mine"]
    account.quarterly_meeting_date = dt.date(2026, 11, 20)
    account.save()

    projected = project_auto_tasks(
        Ability(staff["csm"]).accounts(), TODAY, dt.date(2026, 11, 30), as_of=TODAY
    )
    assert [(p.task_type, p.due_date) for p in projected] == [
        (TaskType.AUTO_QUARTERLY_MEETING, dt.date(2026, 11, 13))
    ]


@time_machine.travel(dt.datetime(2026, 11, 2, 6, 0, tzinfo=UTC))
def test_the_projected_reminder_is_the_one_the_generator_creates(staff, book):
    from csm.tasks.generators import generate_renewal_tasks

    account = book["mine"]
    account.contract_end_date = dt.date(2027, 1, 31)
    account.save()
    day_before = dt.date(2026, 11, 1)
    projected = project_auto_tasks(
        Ability(staff["csm"]).accounts(), day_before, dt.date(2026, 11, 2), as_of=day_before
    )

    generate_renewal_tasks(as_of=dt.date(2026, 11, 2))

    created = Task.objects.get(account=account, task_type=TaskType.AUTO_RENEWAL)
    assert [(p.auto_key, p.due_date, p.title) for p in projected] == [
        (created.auto_key, created.due_date, created.title)
    ]


# ---------------------------------------------------------------------------
# Health
# ---------------------------------------------------------------------------


def history(account, on, status, score):
    return HealthScoreHistory.objects.create(
        account=account,
        calculated_on=on,
        trigger=ScoreTrigger.DAILY,
        score=Decimal(score),
        max_score=100,
        computed_status=status,
        effective_status=status,
        breakdown={
            "factors": {
                "CSAT": {"outcome": "SCORED", "color": "GREEN", "points": "20.50"},
                "OPEN_TICKETS": {"outcome": "OMITTED", "color": None, "points": 0},
            }
        },
    )


def test_health_opens_on_the_last_score_before_the_period(staff, book):
    account = book["mine"]
    history(account, dt.date(2026, 8, 31), HealthStatus.AT_RISK, "45.00")
    history(account, dt.date(2026, 9, 2), HealthStatus.WATCH, "70.00")
    history(account, dt.date(2026, 9, 20), HealthStatus.HEALTHY, "90.00")
    history(book["theirs"], dt.date(2026, 9, 2), HealthStatus.WATCH, "70.00")

    sheets = builders.health(context(staff["csm"], dt.date(2026, 9, 1), dt.date(2026, 9, 30)))

    summary = sheet(sheets, "Summary by account")
    assert [row[:11] for row in summary.rows] == [
        ["REP01", "Mine Co", "Rana CSM", "At Risk", 45.0, "Healthy", 90.0, 45.0, 70.0, 90.0, 2]
    ]
    changes = sheet(sheets, "Status changes")
    assert [(r[3], r[4]) for r in changes.rows] == [("At Risk", "Watch"), ("Watch", "Healthy")]

    calculations = sheet(sheets, "Every calculation")
    assert len(calculations.rows) == 2
    assert column(calculations, "CSAT points") == [20.5, 20.5]
    assert column(calculations, "Open Tickets colour") == ["Not scored", "Not scored"]


def test_health_exports_the_record_behind_each_score_not_only_points(staff, book):
    HealthScoreHistory.objects.create(
        account=book["mine"],
        calculated_on=dt.date(2026, 9, 15),
        trigger=ScoreTrigger.DAILY,
        score=Decimal("70.00"),
        max_score=100,
        computed_status=HealthStatus.WATCH,
        effective_status=HealthStatus.WATCH,
        breakdown={
            "factors": {
                "OPEN_TICKETS": {
                    "outcome": "SCORED", "color": "YELLOW", "points": 50,
                    "detail": {"bugs": 1, "general": 2, "weighted": 5},
                },
                "CSAT": {
                    "outcome": "SCORED", "color": "YELLOW", "points": "60.00",
                    "detail": {
                        "quarter": "2026 Q3", "average": "3.40", "contacts": 2,
                        "ratings": 3, "recorded_on": "2026-09-02",
                    },
                },
                "GENERAL_CONTACT": {
                    "outcome": "SCORED", "color": "RED", "points": 0,
                    "detail": {"days_since": 40, "last_contact_on": "2026-08-06"},
                },
                "MEETING_ATTENDANCE": {
                    "outcome": "OMITTED", "color": None, "points": None,
                    "reason": "required_meeting_frequency_not_set",
                    "detail": {"meetings_in_window": 1},
                },
            }
        },
    )

    sheets = builders.health(context(staff["csm"], *SEPT_RANGE))
    for title in ("Every calculation", "Summary by account"):
        found = sheet(sheets, title)
        record = dict(zip(found.headers, found.rows[0], strict=True))
        assert record["Open bug tickets"] == 1
        assert record["Open general tickets"] == 2
        assert record["CSAT average"] == 3.4
        assert record["CSAT quarter"] == "2026 Q3"
        assert record["CSAT last rated"] == dt.date(2026, 9, 2)
        assert record["Last client contact"] == dt.date(2026, 8, 6)
        assert record["Days since client contact"] == 40
        assert record["Meetings this month"] == 1
        assert record["Meeting Attendance colour"] == "Not scored"
        assert record["Meeting Attendance note"] == "Not scored - no meeting frequency set"
        # A factor this account's set does not include stays blank, not zero.
        assert record["Last non-meeting contact"] is None
        assert record["Weekly Contact colour"] is None


SEPT_RANGE = (dt.date(2026, 9, 1), dt.date(2026, 9, 30))


def test_health_lists_overrides_set_in_the_period(staff, book):
    HealthStatusOverride.objects.create(
        account=book["mine"],
        status=HealthStatus.WATCH,
        reason="Champion left, renewal at risk.",
        score_at_override=Decimal("88.00"),
        max_at_override=100,
        created_by=staff["csm"],
    )
    from csm.core import clock

    # created_at is stamped by the real clock, so the period is the real day.
    day = clock.today()
    overrides = sheet(builders.health(context(staff["csm"], day, day, today=day)), "Overrides")
    assert column(overrides, "Reason") == ["Champion left, renewal at risk."]

    earlier = sheet(
        builders.health(context(staff["csm"], day - dt.timedelta(days=9), day - dt.timedelta(days=1))),
        "Overrides",
    )
    assert earlier.rows == []


# ---------------------------------------------------------------------------
# Renewals and portfolio
# ---------------------------------------------------------------------------


def test_renewals_list_contracts_ending_in_the_period(staff, book):
    book["mine"].contract_end_date = dt.date(2026, 12, 1)
    book["mine"].save()
    book["theirs"].contract_end_date = dt.date(2026, 12, 2)
    book["theirs"].save()

    renewals = sheet(
        builders.renewals(context(staff["senior"], dt.date(2026, 11, 1), dt.date(2026, 12, 31))),
        "Renewals",
    )
    assert column(renewals, "Account") == ["Mine Co", "Theirs Co"]
    assert column(renewals, "Days left") == [44, 45]
    assert "Health status" in renewals.headers


def test_portfolio_has_one_line_per_csm(staff, book):
    done(book["mine"], "One", dt.date(2026, 9, 10), dt.datetime(2026, 9, 12, 8, tzinfo=UTC))
    ActivityFactory(
        account=book["theirs"], occurred_on=dt.date(2026, 9, 3), created_by=staff["other"],
        activity_type=ActivityType.MEETING,
    )

    by_csm = sheet(
        builders.portfolio(context(staff["manager"], dt.date(2026, 9, 1), dt.date(2026, 9, 30))),
        "By CSM",
    )
    rows = {row[0]: dict(zip(by_csm.headers, row, strict=True)) for row in by_csm.rows}
    assert set(rows) == {"Rana CSM", "Omar CSM"}
    assert rows["Rana CSM"]["Tasks completed"] == 1
    assert rows["Rana CSM"]["Completed late"] == 1
    assert rows["Omar CSM"]["Meetings held"] == 1
    assert rows["Omar CSM"]["Client contacts logged"] == 1

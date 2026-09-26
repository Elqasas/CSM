"""Downloadable exports of the accounts and contacts a reader may see.

Headers reuse the importer's own column tables where a field appears in both,
so an export and the sheet it might be re-imported through can't drift into
naming the same column two different ways.

Row scope is the caller's job, same as everywhere else: this module only knows
how to render whatever queryset it is handed.
"""

from __future__ import annotations

import io

from csm.importer.contact_parsing import CONTACT_COLUMNS
from csm.importer.parsing import (
    COLUMNS,
    CUSTOM_PREFIX,
    account_columns,
)
from csm.importer.template import header_label

#: Every built-in field the accounts importer reads, in the same order as
#: ``parsing.COLUMNS`` - which is the add/edit form's order. Not a database
#: dump, so the cached health projection, provenance and generated columns
#: still have no place here. The active custom fields join these at export
#: time, before Description, exactly where the form shows them. Same fields,
#: same order, so a straight re-import of an unedited export is a no-op.
ACCOUNT_FIELDS = list(COLUMNS)

#: Appended only for a reader who holds health.view - never unconditionally.
HEALTH_FIELDS = ["health_status", "health_score"]

#: Every field the contacts importer reads, in the same order as
#: ``CONTACT_COLUMNS`` - same reasoning as ``ACCOUNT_FIELDS`` above.
CONTACT_FIELDS = list(CONTACT_COLUMNS)

#: The two health columns, which only ever appear on the accounts export - not
#: an import field anywhere, so they carry no place in either column table.
_LABELS = {
    "health_status": "Health Status",
    "health_score": "Health Score",
}


def _account_label(field: str, custom_fields=()) -> str:
    if field in COLUMNS or field.startswith(CUSTOM_PREFIX):
        return header_label(field, COLUMNS, custom_fields)
    return _LABELS.get(field, field.replace("_", " ").title())


def _contact_label(field: str) -> str:
    if field in CONTACT_COLUMNS:
        return header_label(field, CONTACT_COLUMNS)
    return _LABELS.get(field, field.replace("_", " ").title())


def _account_row(account, fields) -> list:
    custom = {
        f"{CUSTOM_PREFIX}{value.definition_id}": value.value
        for value in account.custom_field_values.all()
    }
    row = []
    for field in fields:
        if field.startswith(CUSTOM_PREFIX):
            row.append(custom.get(field, ""))
        elif field == "csm":
            row.append(account.csm.email if account.csm else "")
        elif field == "acm":
            row.append(account.acm.email if account.acm else "")
        elif field == "partner":
            row.append(account.partner.name if account.partner else "")
        elif field == "country":
            row.append(account.country.name if account.country else "")
        elif field == "license_type":
            row.append(account.get_license_type_display())
        elif field == "deployment":
            row.append(account.get_deployment_display())
        elif field == "account_class":
            row.append(account.get_account_class_display())
        elif field == "size":
            row.append(account.get_size_display())
        elif field == "lifecycle_stage":
            row.append(account.get_lifecycle_stage_display())
        elif field == "required_meeting_frequency":
            row.append(
                account.get_required_meeting_frequency_display()
                if account.required_meeting_frequency
                else ""
            )
        elif field == "g2_review_status":
            row.append(account.get_g2_review_status_display())
        elif field == "gartner_review_status":
            row.append(account.get_gartner_review_status_display())
        elif field == "health_status":
            row.append(account.effective_health_status or "")
        elif field == "health_score":
            row.append(
                str(account.computed_health_score)
                if account.computed_health_score is not None
                else ""
            )
        else:
            row.append(getattr(account, field))
    return row


def _contact_row(contact, fields) -> list:
    row = []
    for field in fields:
        if field == "client_code":
            row.append(contact.account.client_code)
        elif field == "influence":
            row.append(contact.get_influence_display() if contact.influence else "")
        elif field == "interest":
            row.append(contact.get_interest_display() if contact.interest else "")
        elif field == "channels":
            # Joined with a separator the parser's own splitter recognises, so
            # an unedited export re-imports as the same multi-select.
            row.append(
                ", ".join(c.get_channel_display() for c in contact.channels.all())
            )
        elif field == "preferred_channel":
            row.append(
                contact.get_preferred_channel_display() if contact.preferred_channel else ""
            )
        else:
            row.append(getattr(contact, field))
    return row


def _workbook(headers: list, rows: list) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    workbook = Workbook()
    sheet = workbook.active
    bold = Font(bold=True)

    for index, label in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=index, value=label)
        cell.font = bold
        sheet.column_dimensions[cell.column_letter].width = max(14, min(30, len(label) + 4))
    sheet.freeze_panes = "A2"

    for row_index, values in enumerate(rows, start=2):
        for col_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=col_index, value=value)
            # openpyxl stores any string starting with "=" as a formula, so a
            # description reading "=== renewal notes" would reach Excel as a
            # broken formula - or a live one. Free text stays text.
            if isinstance(value, str) and value.startswith("="):
                cell.data_type = "s"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_accounts_xlsx(queryset, *, include_health: bool, custom_fields=()) -> bytes:
    fields = list(account_columns(custom_fields)) + (
        HEALTH_FIELDS if include_health else []
    )
    headers = [_account_label(f, custom_fields) for f in fields]
    rows = [
        _account_row(account, fields)
        for account in queryset.select_related("csm", "acm", "partner")
        .prefetch_related("custom_field_values")
        .order_by("client_name")
    ]
    return _workbook(headers, rows)


def build_contacts_xlsx(queryset) -> bytes:
    headers = [_contact_label(f) for f in CONTACT_FIELDS]
    rows = [
        _contact_row(contact, CONTACT_FIELDS)
        for contact in queryset.select_related("account")
        .prefetch_related("channels")
        .order_by("account__client_name", "full_name")
    ]
    return _workbook(headers, rows)

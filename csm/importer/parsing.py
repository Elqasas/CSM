"""Reading the client spreadsheet.

The parsing is deliberately forgiving about *shape* and strict about *meaning*.
Header spelling, capitalisation, spacing and punctuation vary between exports of
the same sheet, so headers are normalised and matched against aliases. Values
are not treated the same way: an unrecognised country or an unknown CSM email is
an error on the row, never a guess.

A rule from the settled decisions is enforced here rather than left to the
caller:

* **CSM and ACM columns match on email and never create a user.** An unknown
  address is an error. Auto-creating staff accounts from a spreadsheet would
  make the access system's row scope depend on whatever a sales export happened
  to contain.

Focal points are no longer part of this sheet. They are a property of a
contact (``Contact.is_focal_point``), set from the contacts screen or the
contacts import, not guessed at from an account-level name column.

Nothing here writes to the database except to look users and partners up.
"""

from __future__ import annotations

import datetime as dt
import io
import pathlib
import re
from dataclasses import dataclass, field

from django_countries import countries

from csm.core.enums import (
    AccountClass,
    AccountSize,
    CustomFieldType,
    Deployment,
    LicenseType,
    LifecycleStage,
    MeetingFrequency,
    ReviewStatus,
)

TRUE_WORDS = {"yes", "y", "true", "1", "managed", "managed service", "x"}
FALSE_WORDS = {"no", "n", "false", "0", "", "none", "not managed"}


def normalise_header(value) -> str:
    """Fold a header to a comparable key.

    ``"Client Code"``, ``"client_code"`` and ``"CLIENT  CODE:"`` are the same
    column. Exports differ in all three ways, and a strict match would reject a
    sheet for cosmetic reasons.
    """
    text = str(value or "").strip().lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


#: Column aliases. The first entry of each tuple is what the template emits.
#:
#: Ordered exactly as ``accounts.forms.CANDIDATE_FIELDS`` lays out the add/edit
#: form, so the template, the export and the form read in one arrangement.
#: The logo is the one form field absent: an image has no spreadsheet cell.
#: Admin-defined custom fields slot in before "description" at run time - see
#: ``account_columns``.
COLUMNS: dict[str, tuple] = {
    "client_code": ("client code", "code", "client id", "account code"),
    "client_name": ("client name", "name", "account name", "client"),
    "country": ("country",),
    "location": ("location", "city"),
    "license_type": ("license", "licence", "license type", "edition", "product"),
    "deployment": ("cloud on prem", "deployment", "cloud or on prem", "hosting"),
    "is_managed_service": ("management service", "managed service", "managed"),
    "boq_users_count": ("boq users count", "boq users", "boq", "licensed users"),
    "added_users_count": ("added users", "added users count"),
    "contract_start_date": ("start contract date", "contract start", "start date"),
    "contract_end_date": ("end contract date", "contract end", "end date"),
    "renewal_count": ("renewal count", "renewals", "times renewed"),
    "lifecycle_stage": ("lifecycle stage", "stage", "lifecycle"),
    "csm": ("csm email", "csm", "customer success manager"),
    "acm": ("acm email", "acm", "account manager"),
    "partner": ("partner", "reseller"),
    "quarterly_meeting_date": ("quarterly meeting date", "quarterly meeting"),
    "required_meeting_frequency": (
        "required meeting frequency",
        "meeting frequency",
        "meetings per month",
    ),
    "events_count": ("number of events", "events", "events count"),
    "digital_content_count": ("digital content count", "digital content"),
    "event_location": ("event location",),
    "g2_review_status": ("g2 review status", "g2 review", "g2"),
    "g2_review_url": ("g2 review url", "g2 url"),
    "gartner_review_status": ("gartner review status", "gartner review", "gartner"),
    "gartner_review_url": ("gartner review url", "gartner url"),
    "account_class": ("class", "account class", "tier"),
    "size": ("size", "account size"),
    "website_url": ("website", "website url", "url", "web site"),
    "tenant_url": ("tenant url", "tenant"),
    "drive_url": ("drive url", "drive", "google drive"),
    "linkedin_url": ("linkedin", "linkedin url"),
    "description": ("description", "notes"),
}

REQUIRED_COLUMNS = ("client_code", "client_name")

URL_FIELDS = (
    "g2_review_url",
    "gartner_review_url",
    "website_url",
    "tenant_url",
    "drive_url",
    "linkedin_url",
)

#: Prefix of a custom field's column key. The definition's pk follows it, the
#: same naming ``AccountForm`` gives the matching input.
CUSTOM_PREFIX = "custom_"

#: Where a cleaned value came from, for the keys not named after their column.
#: Used to drop a value whose column the sheet never had - see ``parse``.
_SOURCE_COLUMN = {"csm_id": "csm", "acm_id": "acm", "partner_name": "partner"}


def active_custom_fields() -> list:
    """The custom fields the account form currently shows, in its order."""
    from csm.accounts.models import AccountCustomFieldDefinition

    return list(AccountCustomFieldDefinition.objects.filter(is_active=True))


def builtin_aliases() -> set[str]:
    """Every header a built-in column answers to, normalised."""
    return {alias for aliases in COLUMNS.values() for alias in aliases}


def account_columns(custom_fields=()) -> dict[str, tuple]:
    """``COLUMNS`` with one column per custom field, placed before description.

    A custom field whose name folds onto a built-in header is left out rather
    than allowed to shadow it; the definition form refuses such names, so this
    only guards definitions made before that check existed.
    """
    taken = builtin_aliases()
    custom = {}
    for definition in custom_fields:
        alias = normalise_header(definition.name)
        if alias and alias not in taken:
            taken.add(alias)
            custom[f"{CUSTOM_PREFIX}{definition.pk}"] = (alias,)

    columns = {}
    for name, aliases in COLUMNS.items():
        if name == "description":
            columns.update(custom)
        columns[name] = aliases
    return columns


@dataclass
class ParsedRow:
    row_number: int
    client_code: str = ""
    client_name: str = ""
    cleaned: dict = field(default_factory=dict)
    errors: list = field(default_factory=list)
    warnings: list = field(default_factory=list)

    def error(self, column: str, message: str) -> None:
        self.errors.append({"column": column, "message": message})

    def warn(self, column: str, message: str) -> None:
        self.warnings.append({"column": column, "message": message})


class WorkbookError(Exception):
    """The workbook could not be read, or has no usable header row."""


# ---------------------------------------------------------------------------
# Value coercion
# ---------------------------------------------------------------------------


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _integer(value, row: ParsedRow, column: str, *, default: int = 0) -> int:
    text = _text(value)
    if not text:
        return default
    try:
        return int(float(text.replace(",", "")))
    except (TypeError, ValueError):
        row.error(column, f"{text!r} is not a whole number.")
        return default


def _date(value, row: ParsedRow, column: str):
    if value in (None, ""):
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    text = _text(value)
    for pattern in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%d %b %Y", "%d %B %Y"):
        try:
            return dt.datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    row.error(column, f"{text!r} is not a date this importer recognises.")
    return None


def _boolean(value, row: ParsedRow, column: str) -> bool:
    text = _text(value).lower()
    if text in TRUE_WORDS:
        return True
    if text in FALSE_WORDS:
        return False
    row.error(column, f"{text!r} is not a yes or no.")
    return False


def _choice(value, choices, row: ParsedRow, column: str, *, default=None, required=False):
    """Match a cell against an enum by value or by label, case-insensitively."""
    text = _text(value)
    if not text:
        if required:
            row.error(column, "This column is required.")
        return default

    folded = normalise_header(text)
    for stored, label in choices:
        if folded in (normalise_header(stored), normalise_header(label)):
            return stored

    allowed = ", ".join(label for _, label in choices)
    row.error(column, f"{text!r} is not one of: {allowed}.")
    return default


def _url(value, row: ParsedRow, column: str) -> str:
    """A link, with https assumed when the cell omits the scheme - the same
    assumption the account form makes."""
    from django.core.exceptions import ValidationError
    from django.core.validators import URLValidator

    text = _text(value)
    if not text:
        return ""
    if "://" not in text:
        text = f"https://{text}"
    if len(text) > 500:
        row.error(column, "Longer than 500 characters.")
        return ""
    try:
        URLValidator()(text)
    except ValidationError:
        row.error(column, f"{_text(value)!r} is not a web address.")
        return ""
    return text


def _custom_value(value, definition, row: ParsedRow) -> str:
    """One custom field's cell, stored as text exactly as the form stores it."""
    column = definition.name
    if definition.field_type == CustomFieldType.DATETIME:
        if value in (None, ""):
            return ""
        if isinstance(value, dt.datetime):
            return value.isoformat()
        if isinstance(value, dt.date):
            return dt.datetime.combine(value, dt.time()).isoformat()
        text = _text(value)
        try:
            return dt.datetime.fromisoformat(text).isoformat()
        except ValueError:
            pass
        for pattern in ("%d/%m/%Y %H:%M", "%m/%d/%Y %H:%M", "%d/%m/%Y", "%m/%d/%Y"):
            try:
                return dt.datetime.strptime(text, pattern).isoformat()
            except ValueError:
                continue
        row.error(column, f"{text!r} is not a date and time this importer recognises.")
        return ""

    text = _text(value)
    if definition.field_type == CustomFieldType.DROPDOWN and text:
        options = definition.choice_list()
        for option in options:
            if normalise_header(option) == normalise_header(text):
                return option
        row.error(column, f"{text!r} is not one of: {', '.join(options)}.")
        return ""
    return text


def _country(value, row: ParsedRow, column: str) -> str:
    text = _text(value)
    if not text:
        return ""
    if len(text) == 2 and text.upper() in {code for code, _ in countries}:
        return text.upper()
    folded = normalise_header(text)
    for code, name in countries:
        if normalise_header(name) == folded:
            return code
    row.error(column, f"{text!r} is not a country name or two-letter code.")
    return ""


# ---------------------------------------------------------------------------
# Row parsing
# ---------------------------------------------------------------------------


def _resolve_user(email_value, row: ParsedRow, column: str, people: dict):
    """Match a staff member by email. Never create one.

    Email is unique, so there is exactly one match or none - an unknown
    address is an error rather than a guess.
    """
    email_text = _text(email_value)
    if not email_text:
        return None

    if "@" not in email_text:
        row.error(column, f"{email_text!r} is not an email address.")
        return None

    user = people["by_email"].get(email_text.lower())
    if user is None:
        row.error(column, f"No active user has the email {email_text!r}.")
    return user


def parse_row(row_number: int, values: dict, *, people: dict) -> ParsedRow:
    """Turn one spreadsheet row into cleaned values plus errors."""
    row = ParsedRow(row_number=row_number)

    row.client_code = _text(values.get("client_code"))
    row.client_name = _text(values.get("client_name"))
    if not row.client_code:
        row.error("Client code", "This column is required.")
    if not row.client_name:
        row.error("Client name", "This column is required.")

    cleaned: dict = {
        "client_code": row.client_code,
        "client_name": row.client_name,
        "location": _text(values.get("location")),
        "event_location": _text(values.get("event_location")),
        "description": _text(values.get("description")),
    }

    for name in URL_FIELDS:
        cleaned[name] = _url(values.get(name), row, header_label_for(name))

    cleaned["renewal_count"] = _integer(values.get("renewal_count"), row, "Renewal count")
    if cleaned["renewal_count"] < 0:
        row.error("Renewal count", "Cannot be negative.")
        cleaned["renewal_count"] = 0

    cleaned["license_type"] = _choice(
        values.get("license_type"), LicenseType.choices, row, "License", default=""
    )
    cleaned["country"] = _country(values.get("country"), row, "Country")
    cleaned["deployment"] = _choice(
        values.get("deployment"), Deployment.choices, row, "Cloud/on-prem",
        default=Deployment.CLOUD,
    )
    cleaned["is_managed_service"] = _boolean(
        values.get("is_managed_service"), row, "Management service"
    )
    cleaned["boq_users_count"] = _integer(values.get("boq_users_count"), row, "BOQ users")
    cleaned["added_users_count"] = _integer(values.get("added_users_count"), row, "Added users")
    cleaned["events_count"] = _integer(values.get("events_count"), row, "Number of events")
    cleaned["digital_content_count"] = _integer(
        values.get("digital_content_count"), row, "Digital content count"
    )

    cleaned["contract_start_date"] = _date(
        values.get("contract_start_date"), row, "Start contract date"
    )
    cleaned["contract_end_date"] = _date(
        values.get("contract_end_date"), row, "End contract date"
    )
    cleaned["quarterly_meeting_date"] = _date(
        values.get("quarterly_meeting_date"), row, "Quarterly meeting date"
    )

    if (
        cleaned["contract_start_date"]
        and cleaned["contract_end_date"]
        and cleaned["contract_end_date"] < cleaned["contract_start_date"]
    ):
        row.error("End contract date", "The contract ends before it starts.")

    cleaned["account_class"] = _choice(
        values.get("account_class"), AccountClass.choices, row, "Class",
        default=AccountClass.C,
    )
    cleaned["size"] = _choice(
        values.get("size"), AccountSize.choices, row, "Size", default=AccountSize.MID
    )
    cleaned["g2_review_status"] = _choice(
        values.get("g2_review_status"), ReviewStatus.choices, row, "G2 review status",
        default=ReviewStatus.NONE,
    )
    cleaned["gartner_review_status"] = _choice(
        values.get("gartner_review_status"), ReviewStatus.choices, row,
        "Gartner review status", default=ReviewStatus.NONE,
    )

    # Lifecycle stage is optional and defaults to Active, which is what starts
    # the scoring immediately. An empty cell is not an error.
    stage_text = _text(values.get("lifecycle_stage"))
    if stage_text:
        cleaned["lifecycle_stage"] = _choice(
            values.get("lifecycle_stage"), LifecycleStage.choices, row, "Lifecycle stage",
            default=LifecycleStage.ACTIVE,
        )
    else:
        cleaned["lifecycle_stage"] = LifecycleStage.ACTIVE
        row.warn("Lifecycle stage", "Empty, so this account imports as Active.")

    frequency_text = _text(values.get("required_meeting_frequency"))
    if frequency_text:
        cleaned["required_meeting_frequency"] = _choice(
            values.get("required_meeting_frequency"),
            MeetingFrequency.choices, row, "Required meeting frequency",
        )
    else:
        cleaned["required_meeting_frequency"] = None
        if cleaned["is_managed_service"]:
            row.warn(
                "Required meeting frequency",
                "Empty on a managed account, so Meeting Attendance is not scored "
                "and this account is marked out of 80 and capped at Watch.",
            )

    csm = _resolve_user(values.get("csm"), row, "CSM", people)
    acm = _resolve_user(values.get("acm"), row, "ACM", people)
    cleaned["csm_id"] = csm.pk if csm else None
    cleaned["acm_id"] = acm.pk if acm else None
    if csm is None and not _text(values.get("csm")):
        row.warn("CSM", "Empty, so this account imports unassigned.")

    cleaned["partner_name"] = _text(values.get("partner"))

    row.cleaned = cleaned
    return row


# ---------------------------------------------------------------------------
# Workbook reading
# ---------------------------------------------------------------------------


def _rows_from_xlsx(file_obj):
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(file_obj, data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises many shapes
        raise WorkbookError(
            f"This file could not be read as a spreadsheet: {exc}"
        ) from exc

    return list(workbook.worksheets[0].iter_rows(values_only=True))


def _rows_from_csv(file_obj):
    """Read a CSV the way the spreadsheet that produced it actually wrote it.

    Two real-world details rather than defensiveness for its own sake. Excel
    writes a UTF-8 byte-order mark, which would otherwise ride along on the
    first header and stop "Client Code" matching anything. And Excel emits
    semicolon-separated files in locales whose decimal separator is a comma, so
    the delimiter is sniffed rather than assumed - a semicolon file read as
    comma-separated parses as one enormous column and validates as empty.
    """
    import csv

    raw = file_obj.read()
    if isinstance(raw, bytes):
        for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise WorkbookError("This file's text encoding could not be determined.")
    else:
        text = raw

    text = text.lstrip("﻿")
    if not text.strip():
        raise WorkbookError("The file is empty.")

    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;	|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","

    return [tuple(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


#: Extensions the importer will read, and the reader for each.
READERS = {
    ".xlsx": _rows_from_xlsx,
    ".xlsm": _rows_from_xlsx,
    ".csv": _rows_from_csv,
    ".txt": _rows_from_csv,
}


def _is_comment(row) -> bool:
    if not row:
        return False
    first = row[0]
    return isinstance(first, str) and first.lstrip().startswith("#")


def read_table(file_obj, *, filename: str = "", columns: dict | None = None,
               required: tuple = ()) -> tuple:
    """Return ``(header_map, body, unknown)`` from a CSV or worksheet.

    ``header_map`` maps our field names to column indexes. Unrecognised columns
    are reported rather than ignored silently: a column the importer skipped is
    data somebody typed and expects to see.
    """
    columns = columns if columns is not None else COLUMNS
    required = required or REQUIRED_COLUMNS

    suffix = pathlib.PurePath(filename or "").suffix.lower()
    reader = READERS.get(suffix, _rows_from_xlsx)
    all_rows = reader(file_obj)

    # A row whose first cell starts with "#" is commentary, not data. The CSV
    # templates carry their guidance that way, having nowhere to put an
    # instructions sheet, and the example row is commented out for the same
    # reason: a template must not import a fictional client on first upload.
    all_rows = [row for row in all_rows if not _is_comment(row)]

    rows = iter(all_rows)
    try:
        header = next(rows)
    except StopIteration:
        raise WorkbookError("The file has no header row.") from None

    alias_to_field = {}
    for field_name, aliases in columns.items():
        for alias in aliases:
            alias_to_field[alias] = field_name

    header_map: dict[str, int] = {}
    unknown: list[str] = []
    for index, cell in enumerate(header):
        key = normalise_header(cell)
        if not key:
            continue
        field_name = alias_to_field.get(key)
        if field_name is None:
            unknown.append(str(cell).strip())
        elif field_name not in header_map:
            header_map[field_name] = index

    missing = [name for name in required if name not in header_map]
    if missing:
        wanted = ", ".join(columns[name][0].title() for name in missing)
        raise WorkbookError(
            f"The sheet is missing required column(s): {wanted}. "
            "Download the template to see the expected headers."
        )

    body = []
    for offset, values in enumerate(rows, start=2):
        if values is None or all(v in (None, "") for v in values):
            continue  # a blank spacer row is not an error
        body.append((offset, values))

    return header_map, body, unknown


def read_workbook(file_obj) -> tuple:
    """Backwards-compatible alias for an xlsx read."""
    return read_table(file_obj, filename="sheet.xlsx")


def header_label_for(field_name: str) -> str:
    """The column's display name, for error messages."""
    from csm.importer.template import header_label

    return header_label(field_name)


def _keep_present(cleaned: dict, present) -> dict:
    """Drop every value whose column the sheet does not have.

    An update writes what the row carries, so a value defaulted for a missing
    column would overwrite real data - a sheet exported before Description
    existed would blank every account's description on re-import. A column that
    is present but empty still clears the field: that is somebody's edit.
    """
    return {
        key: value
        for key, value in cleaned.items()
        if key in REQUIRED_COLUMNS or _SOURCE_COLUMN.get(key, key) in present
    }


def parse(file_obj, *, people: dict, filename: str = "",
          custom_fields=None) -> tuple:
    """Parse an accounts sheet into ``(ParsedRow list, unknown column list)``.

    ``custom_fields`` are the active definitions to recognise as columns; the
    caller passes them, like ``people``, so this module reads nothing itself.
    """
    custom_fields = list(custom_fields or [])
    definitions = {f"{CUSTOM_PREFIX}{d.pk}": d for d in custom_fields}
    header_map, body, unknown = read_table(
        file_obj, filename=filename, columns=account_columns(custom_fields)
    )

    parsed = []
    seen_codes: dict[str, int] = {}
    for row_number, values in body:
        picked = {
            name: (values[index] if index < len(values) else None)
            for name, index in header_map.items()
        }
        row = parse_row(row_number, picked, people=people)
        row.cleaned = _keep_present(row.cleaned, picked)

        custom_values = {
            str(definitions[key].pk): _custom_value(value, definitions[key], row)
            for key, value in picked.items()
            if key in definitions
        }
        if custom_values:
            row.cleaned["custom_fields"] = custom_values

        # A code repeated inside one sheet is almost always two versions of the
        # same client rather than two clients, and applying both would leave
        # whichever sorted last silently overwriting the other.
        if row.client_code:
            first = seen_codes.get(row.client_code.lower())
            if first is not None:
                row.error(
                    "Client code",
                    f"{row.client_code!r} also appears on row {first} of this sheet.",
                )
            else:
                seen_codes[row.client_code.lower()] = row_number

        parsed.append(row)

    return parsed, unknown

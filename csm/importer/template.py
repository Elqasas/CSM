"""The downloadable import template.

Headers come from ``parsing.COLUMNS`` rather than being typed out again, so the
template a person fills in cannot drift from the aliases the parser accepts.
That drift is the failure this file exists to prevent: a template with a header
the importer no longer recognises produces a sheet that validates as empty.
"""

from __future__ import annotations

import io

from csm.core.enums import (
    AccountClass,
    AccountSize,
    Deployment,
    LifecycleStage,
    MeetingFrequency,
    ReviewStatus,
)
from csm.importer.parsing import COLUMNS, REQUIRED_COLUMNS

#: A short note per column, written into the second row.
GUIDANCE = {
    "license_type": (
        "The product edition: Advantage / Complete / Drill / Discover / "
        "Essential / Phishing."
    ),
    "boq_users_count": "Seats contracted for. A whole number, not the edition.",
    "client_code": "Required. The natural key; rows match on it.",
    "client_name": "Required.",
    "country": "Country name or two-letter code.",
    "deployment": " / ".join(label for _, label in Deployment.choices),
    "is_managed_service": "Yes or No.",
    "csm": "dPhish email address of the CSM.",
    "acm": "dPhish email address of the ACM.",
    "partner": "Created if it does not exist yet.",
    "contract_start_date": "YYYY-MM-DD.",
    "contract_end_date": "YYYY-MM-DD.",
    "quarterly_meeting_date": "YYYY-MM-DD.",
    "account_class": " / ".join(label for _, label in AccountClass.choices),
    "size": " / ".join(label for _, label in AccountSize.choices),
    "lifecycle_stage": "Optional. Empty means Active. "
                       + " / ".join(label for _, label in LifecycleStage.choices),
    "required_meeting_frequency": "Optional, managed accounts only. "
                                  + " / ".join(label for _, label in MeetingFrequency.choices),
    "g2_review_status": " / ".join(label for _, label in ReviewStatus.choices),
    "gartner_review_status": " / ".join(label for _, label in ReviewStatus.choices),
    "renewal_count": "How many times this client has renewed. A whole number.",
    "g2_review_url": "Link to the published G2 review.",
    "gartner_review_url": "Link to the published Gartner review.",
    "website_url": "The client's website. https:// is added if left off.",
    "tenant_url": "Link to the client's tenant.",
    "drive_url": "Link to the account's Drive folder.",
    "linkedin_url": "The client's LinkedIn page.",
    "description": "Free-text notes about the account.",
}


def custom_guidance(definition) -> str:
    """The note for an admin-defined column, which GUIDANCE cannot know."""
    from csm.core.enums import CustomFieldType

    if definition.field_type == CustomFieldType.DROPDOWN:
        return "Custom field. One of: " + " / ".join(definition.choice_list()) + "."
    if definition.field_type == CustomFieldType.DATETIME:
        return "Custom field. Date and time, YYYY-MM-DD HH:MM."
    return "Custom field. Free text."


def account_guidance(custom_fields=()) -> dict:
    """GUIDANCE plus a note for each custom-field column."""
    from csm.importer.parsing import CUSTOM_PREFIX

    guidance = dict(GUIDANCE)
    for definition in custom_fields:
        guidance[f"{CUSTOM_PREFIX}{definition.pk}"] = custom_guidance(definition)
    return guidance


def _contact_guidance() -> dict:
    from csm.importer.csv_template import CONTACT_GUIDANCE

    return CONTACT_GUIDANCE


#: Columns whose canonical alias does not title-case into something a person
#: would write. The parser normalises headers anyway, so these are cosmetic -
#: but a template is the first thing anyone sees and "Boq Users Count" reads as
#: though nobody looked at it.
LABEL_OVERRIDES = {
    "boq_users_count": "BOQ Users Count",
    "license_type": "License",
    "deployment": "Cloud / On-prem",
    "csm": "CSM Email",
    "acm": "ACM Email",
    "g2_review_status": "G2 Review Status",
    "g2_review_url": "G2 Review URL",
    "gartner_review_url": "Gartner Review URL",
    "tenant_url": "Tenant URL",
    "drive_url": "Drive URL",
    "linkedin_url": "LinkedIn",
    "is_managed_service": "Management Service",
    "email": "Email",
}


def header_label(field_name: str, columns: dict | None = None,
                 custom_fields=()) -> str:
    """The canonical spelling: an override, or the first alias title-cased.

    A custom field's column is headed with its name exactly as the admin wrote
    it, which is also what the account form labels it.
    """
    from csm.importer.parsing import CUSTOM_PREFIX

    if field_name.startswith(CUSTOM_PREFIX):
        for definition in custom_fields:
            if f"{CUSTOM_PREFIX}{definition.pk}" == field_name:
                return definition.name
    if field_name in LABEL_OVERRIDES:
        return LABEL_OVERRIDES[field_name]
    return (columns or COLUMNS)[field_name][0].title()


def build_template(columns: dict | None = None, required: tuple | None = None,
                   *, custom_fields=()) -> bytes:
    """A headers-only first sheet, with the guidance on a second one.

    The guidance deliberately does not sit under the headers. The parser reads
    the first worksheet and treats every non-empty row after the header as data,
    so a row of help text would import as an account called "Required. The
    natural key" - the template would poison its own import.

    ``custom_fields`` are the active definitions, passed by the caller so this
    builder reads no database; each becomes a column before Description.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    from csm.importer.parsing import account_columns

    is_accounts = columns is None or columns is COLUMNS
    custom_fields = list(custom_fields) if is_accounts else []
    if is_accounts:
        columns = account_columns(custom_fields)
        guidance = account_guidance(custom_fields)
    else:
        guidance = _contact_guidance()
    required = required if required is not None else REQUIRED_COLUMNS

    def label_of(field_name):
        return header_label(field_name, columns, custom_fields)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Accounts" if is_accounts else "Contacts"

    bold = Font(bold=True)

    for index, field_name in enumerate(columns, start=1):
        label = label_of(field_name)
        if field_name in required:
            label = f"{label} *"
        cell = sheet.cell(row=1, column=index, value=label)
        cell.font = bold
        sheet.column_dimensions[cell.column_letter].width = max(
            16, min(34, len(label) + 4)
        )

    sheet.freeze_panes = "A2"

    guide = workbook.create_sheet("Instructions")
    guide.cell(row=1, column=1, value="Column").font = bold
    guide.cell(row=1, column=2, value="Notes").font = bold
    guide.column_dimensions["A"].width = 30
    guide.column_dimensions["B"].width = 80

    for offset, field_name in enumerate(columns, start=2):
        guide.cell(row=offset, column=1, value=label_of(field_name))
        note = guide.cell(row=offset, column=2, value=guidance.get(field_name, ""))
        note.alignment = Alignment(wrap_text=True, vertical="top")

    guide.cell(
        row=len(columns) + 3,
        column=1,
        value="Rows with errors are skipped; the rest of the sheet still imports.",
    )

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()

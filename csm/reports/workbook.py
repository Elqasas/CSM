"""Turning report sheets into an .xlsx file.

One writer for every download the platform offers - the reports and the
accounts and contacts sheets alike - so the guard against a cell turning into a
live formula is written once.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime

from csm.core.clock import business_tz


@dataclass
class Sheet:
    title: str
    headers: list[str]
    rows: list[list] = field(default_factory=list)


#: Excel refuses a worksheet title longer than this.
_TITLE_LIMIT = 31


def build_workbook(sheets: list[Sheet]) -> bytes:
    """Write ``sheets`` in order, each with a bold frozen header row.

    Write-only mode, so a two-year health history streams to the file rather
    than holding every cell as an object at once.
    """
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook(write_only=True)
    bold = Font(bold=True)

    for sheet in sheets:
        worksheet = workbook.create_sheet(title=sheet.title[:_TITLE_LIMIT])
        for index, label in enumerate(sheet.headers, start=1):
            worksheet.column_dimensions[get_column_letter(index)].width = max(
                12, min(40, len(label) + 4)
            )
        worksheet.freeze_panes = "A2"

        header = []
        for label in sheet.headers:
            cell = WriteOnlyCell(worksheet, value=label)
            cell.font = bold
            header.append(cell)
        worksheet.append(header)

        for values in sheet.rows:
            worksheet.append([_cell(worksheet, value, WriteOnlyCell) for value in values])

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _cell(worksheet, value, cell_class):
    if isinstance(value, str) and value.startswith("="):
        # openpyxl stores any string starting with "=" as a formula, so a note
        # reading "=== renewal notes" would reach Excel as a broken formula - or
        # a live one. Free text stays text.
        cell = cell_class(worksheet, value=value)
        cell.data_type = "s"
        return cell
    if isinstance(value, datetime):
        # Excel has no timezones. A stored instant is UTC; the reader thinks in
        # Cairo, so it is written as the Cairo wall-clock time.
        if value.tzinfo is not None:
            value = value.astimezone(business_tz())
        cell = cell_class(worksheet, value=value.replace(tzinfo=None))
        cell.number_format = "yyyy-mm-dd hh:mm"
        return cell
    if isinstance(value, date):
        cell = cell_class(worksheet, value=value)
        cell.number_format = "yyyy-mm-dd"
        return cell
    return value
